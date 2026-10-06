import docx

doc = docx.Document("BAB 1_1438.docx")
start_dp = -1
for i, p in enumerate(doc.paragraphs):
    if "DAFTAR PUSTAKA" in p.text.upper():
        start_dp = i
        break

dp_entries = []
for i in range(start_dp + 1, len(doc.paragraphs)):
    t = doc.paragraphs[i].text.strip()
    if t:
        dp_entries.append((i, t))

print(f"Total entries in BAB 1_1438.docx DAFTAR PUSTAKA: {len(dp_entries)}")
for idx, (p_num, text) in enumerate(dp_entries, 1):
    author = text.split("(")[0].strip() if "(" in text else text[:30]
    year = text.split("(")[1].split(")")[0] if "(" in text and ")" in text else ""
    print(f"{idx:2d}. [{p_num}] {author} ({year})")

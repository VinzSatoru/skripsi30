import docx

with open("BAB_1_1438_temp.docx", "rb") as f:
    doc = docx.Document(f)

print("=== PARAGRAPHS 31 to 65 ===")
for i in range(31, min(65, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    print(f"[{i}] ({p.style.name}): {p.text}")

print("\n=== TABLES IN DOC ===")
for i, table in enumerate(doc.tables):
    # Print first cell or row text
    sample = " | ".join([cell.text.strip().replace('\n', ' ') for cell in table.rows[0].cells[:3]])
    print(f"Table {i}: rows={len(table.rows)}, cols={len(table.columns)}, Header sample: {sample[:100]}")

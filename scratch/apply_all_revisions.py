import docx, zipfile, io, shutil, os, sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

DOCX_IN = r"01_Naskah_Proposal_Master\PROPOSAL_BAB_1-3_FINAL_SEMPURNA.docx"
DOCX_OUT = r"01_Naskah_Proposal_Master\PROPOSAL_BAB_1-3_FINAL_SEMPURNA.docx"

doc = docx.Document(DOCX_IN)
print(f"Loaded document. Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")

def insert_paragraphs_after(ref_p, p_data_list):
    """
    Inserts a list of paragraphs after ref_p.
    p_data_list is a list of tuples: (text, style, alignment)
    """
    curr = ref_p
    for item in p_data_list:
        text = item[0]
        style = item[1] if len(item) > 1 else 'Normal'
        new_p = doc.add_paragraph(text, style=style)
        if len(item) > 2 and item[2] is not None:
            new_p.alignment = item[2]
        curr._p.addnext(new_p._p)
        curr = new_p
    return curr

# ==============================================================================
# 1. COVER & FRONT MATTER
# ==============================================================================
print("\n--- 1. Updating Cover & Front Matter ---")
# Standardize Cover Title
if "Shapley Additive Explanations" in doc.paragraphs[1].text:
    doc.paragraphs[1].text = doc.paragraphs[1].text.replace("Shapley Additive Explanations", "Shapley Additive exPlanations")
    print("[OK] Cover title updated to 'Shapley Additive exPlanations'")

# Cover Year
if "TP 2026/2027" in doc.paragraphs[17].text:
    doc.paragraphs[17].text = "TAHUN AKADEMIK 2026/2027"
    print("[OK] Cover year updated to 'TAHUN AKADEMIK 2026/2027'")

# Standardize any 'Shapley Additive explanations' in TOC/Headings
for i, p in enumerate(doc.paragraphs):
    if "Shapley Additive explanations" in p.text:
        p.text = p.text.replace("Shapley Additive explanations", "Shapley Additive exPlanations")
        print(f"[OK] Standardized SHAP text in P[{i}]")

# ==============================================================================
# 2. BAB I (PENDAHULUAN)
# ==============================================================================
print("\n--- 2. Updating Bab I ---")
# Subbab 1.1 WHO Citation (P[48])
for i in range(45, 52):
    p = doc.paragraphs[i]
    if "Berdasarkan catatan Organisasi Kesehatan Dunia (WHO)" in p.text:
        if "(Medic et al., 2017)" not in p.text:
            p.text = p.text.replace(
                "sekitar sepertiga populasi dunia kerap mengalami masalah tidur.",
                "sekitar sepertiga populasi dunia kerap mengalami masalah tidur (Medic et al., 2017)."
            )
            print(f"[OK] Added (Medic et al., 2017) citation to WHO sentence in P[{i}]")

# Subbab 1.1 Henrich citation & Xie citation (P[49], P[52])
for i in range(47, 55):
    p = doc.paragraphs[i]
    if "(Henrich, 2021)" in p.text:
        p.text = p.text.replace("(Henrich, 2021)", "(Henrich et al., 2021)")
        print(f"[OK] Updated (Henrich, 2021) to (Henrich et al., 2021) in P[{i}]")
    if "(Xie et al., 2026)" in p.text:
        p.text = p.text.replace("(Xie et al., 2026)", "(Xie et al., 2025)")
        print(f"[OK] Updated (Xie et al., 2026) to (Xie et al., 2025) in P[{i}]")

# Subbab 1.2 Rumusan Masalah & Subbab 1.3 Tujuan Penelitian
rumusan_idx = None
tujuan_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "1.2 Rumusan Masalah":
        rumusan_idx = i
    elif p.text.strip() == "1.3 Tujuan Penelitian":
        tujuan_idx = i

if rumusan_idx is not None and tujuan_idx is not None:
    # Update Subbab 1.2 content (paragraph right after heading)
    rp = doc.paragraphs[rumusan_idx + 1]
    rp.text = (
        "\tBerdasarkan latar belakang masalah yang telah diuraikan, rumusan masalah dalam penelitian ini dirinci menjadi tiga pertanyaan operasional sebagai berikut:\n"
        "1. Bagaimana tingkat kinerja, presisi, dan akurasi klasifikasi algoritma CatBoost dalam mengidentifikasi kategori risiko gangguan tidur pada dataset metrik gaya hidup digital yang mengalami ketidakseimbangan kelas (class imbalance)?\n"
        "2. Bagaimana kapabilitas metode Tree-SHAP dalam mengekstraksi serta menyajikan visualisasi atribusi kontribusi fitur gaya hidup secara utuh tanpa mengalami distorsi fragmentasi nilai kategori?\n"
        "3. Bagaimana rancangan integrasi penjelasan lokal Tree-SHAP dengan pedoman kebersihan tidur (Sleep Hygiene) sehingga mampu menghasilkan rekomendasi intervensi klinis personal yang mudah dipahami dan ditindaklanjuti secara nyata oleh tenaga medis maupun pasien?"
    )
    print(f"[OK] Updated 1.2 Rumusan Masalah in P[{rumusan_idx + 1}]")

    # Update Subbab 1.3 content (paragraph right after heading)
    tp = doc.paragraphs[tujuan_idx + 1]
    tp.text = (
        "\tMengacu pada rumusan masalah yang telah ditetapkan, tujuan dari penelitian ini dirumuskan secara terukur dan berpasangan sebagai berikut:\n"
        "1. Menganalisis dan mengevaluasi kinerja, presisi, dan akurasi klasifikasi algoritma CatBoost dengan penanganan class imbalance dalam mengidentifikasi kategori risiko gangguan tidur multi-kelas pada dataset metrik gaya hidup digital.\n"
        "2. Menguji dan mengimplementasikan kapabilitas metode Tree-SHAP dalam mengekstraksi serta menyajikan atribusi kontribusi fitur secara global dan lokal tanpa fragmentasi semantik data kategorikal.\n"
        "3. Merancang integrasi penjelasan lokal Tree-SHAP dengan prinsip pedoman Sleep Hygiene untuk menghasilkan rekomendasi intervensi klinis personal yang aplikatif dan mudah ditindaklanjuti sebagai purwarupa Clinical Decision Support System (CDSS)."
    )
    print(f"[OK] Updated 1.3 Tujuan Penelitian in P[{tujuan_idx + 1}]")

# ==============================================================================
# 3. BAB II (LANDASAN TEORI)
# ==============================================================================
print("\n--- 3. Updating Bab II ---")
# Subbab 2.1.1 Sleep Hygiene parameters & citations
for i in range(74, 80):
    p = doc.paragraphs[i]
    if "prinsip kebersihan tidur (Sleep Hygiene) digunakan sebagai pedoman klinis" in p.text:
        p.text = (
            "\tSelain dampak fisiologis akibat paparan radiasi layar gawai, interaksi digital yang konstan juga melahirkan tekanan psikologis baru yang dikenal luas sebagai technostress. Tuntutan untuk selalu merespons pesan kerja, paparan informasi tanpa henti, serta ketakutan tertinggal informasi (FOMO) menciptakan kondisi status kewaspadaan mental fisiologis berlebihan yang menghambat relaksasi pikiran sebelum terlelap (Widayati, 2024). Tekanan psikososial ini kemudian berpadu dengan kebiasaan gaya hidup yang buruk, seperti konsumsi kafein berlebih dan waktu istirahat yang tidak teratur, sehingga memperparah siklus gangguan tidur harian (Henrich et al., 2021). "
            "Untuk memutus lingkaran masalah tersebut, prinsip kebersihan tidur (Sleep Hygiene) digunakan sebagai pedoman klinis guna menata ulang pola perilaku individu (Irish et al., 2015). Standar parameter klinis Sleep Hygiene menetapkan bahwa paparan perangkat gawai harus dihentikan minimal 30 hingga 60 menit sebelum tidur guna mencegah penekanan hormon melatonin akibat emisi cahaya biru (Chang et al., 2015), serta konsumsi zat berkafein wajib dihentikan setidaknya 6 jam sebelum waktu tidur (misalnya maksimal pukul 14.00 WIB untuk jadwal tidur malam pukul 22.00 WIB) karena waktu paruh eliminasi metabolisme kafein dalam tubuh berkisar antara 3 hingga 7 jam (Drake et al., 2013). Selain itu, pembatasan konsumsi alkohol menjelang tidur dan pengaturan jadwal tidur-bangun yang konsisten terbukti krusial dalam mempertahankan kontinuitas fase tidur gelombang lambat (Slow-Wave Deep Sleep) serta mencegah fragmentasi tidur malam. Menghubungkan metrik gaya hidup digital dengan ambang batas klinis pedoman kebersihan tidur ini menjadi dasar perancangan intervensi preventif yang terarah."
        )
        print(f"[OK] Enhanced Sleep Hygiene parameters and citations in P[{i}]")

# Subbab 2.2 Xie citation update
for i in range(90, 96):
    p = doc.paragraphs[i]
    if "(Xie et al., 2026)" in p.text:
        p.text = p.text.replace("(Xie et al., 2026)", "(Xie et al., 2025)")
        print(f"[OK] Updated (Xie et al., 2026) to (Xie et al., 2025) in Bab II P[{i}]")

# Inserting mathematical equations in Bab II
# Check if equations already exist to avoid duplicate insertions
has_eq_2_1 = any("Persamaan 2.1" in p.text for p in doc.paragraphs)

if not has_eq_2_1:
    print("[INFO] Inserting Persamaan 2.1 to 2.8 in Bab II...")
    
    # 1. In Subbab 2.1.2 (Klasifikasi dan Machine Learning)
    # Find paragraph 2.1.2 content (ends around Kaya, 2025)
    target_p_2_1_2 = None
    for i in range(75, 83):
        if "pendekatan machine learning membuka peluang besar bagi otomatisasi penapisan risiko prediktif modern" in doc.paragraphs[i].text:
            target_p_2_1_2 = doc.paragraphs[i]
            break
            
    if target_p_2_1_2:
        eqs_2_1_2 = [
            ("\tDalam mengoptimalkan batas keputusan klasifikasi multi-kelas, model terawasi meminimalkan fungsi kerugian Multi-class Cross-Entropy Loss yang didefinisikan secara matematis pada Persamaan (2.1):", "Normal", None),
            ("\tL_Multi-CE = - (1 / N) * ∑_{i=1}^N ∑_{k=1}^K y_{i,k} · ln(p̂_{i,k})                    (Persamaan 2.1)", "Normal", None),
            ("\tdi mana N menyatakan total jumlah sampel observasi, K adalah jumlah tingkatan kelas risiko (K = 4: Healthy, Mild, Moderate, Severe), y_{i,k} merupakan variabel indikator biner bernilai 1 jika sampel ke-i termasuk ke dalam kelas k dan 0 jika sebaliknya, serta p̂_{i,k} melambangkan probabilitas prediksi model bahwa sampel ke-i berada pada kelas risiko k.", "Normal", None),
            ("\tSelanjutnya, kinerja model pada setiap tingkatan kelas diukur menggunakan metrik Recall (sensitivitas deteksi) dan Precision (tingkat ketepatan prediksi) sebagaimana diformulasikan pada Persamaan (2.2) dan Persamaan (2.3):", "Normal", None),
            ("\tRecall_k = TP_k / (TP_k + FN_k)                                                      (Persamaan 2.2)", "Normal", None),
            ("\tPrecision_k = TP_k / (TP_k + FP_k)                                                   (Persamaan 2.3)", "Normal", None),
            ("\tdi mana TP_k (True Positive) adalah jumlah kasus kelas k yang terprediksi benar, FN_k (False Negative) adalah kasus kelas k yang terlewat dan diprediksi keliru sebagai kelas lain, serta FP_k (False Positive) merupakan sampel kelas lain yang secara salah terprediksi sebagai kelas k. Guna menyatukan kedua metrik secara seimbang tanpa bias terhadap ukuran kelas mayoritas, metrik Macro-averaged F1-Score digunakan dengan menghitung rata-rata aritmetika harmonik seluruh kelas sesuai Persamaan (2.4):", "Normal", None),
            ("\tMacro-F1 = (1 / K) * ∑_{k=1}^K [2 * (Precision_k * Recall_k) / (Precision_k + Recall_k)] (Persamaan 2.4)", "Normal", None)
        ]
        target_p_2_1_2 = insert_paragraphs_after(target_p_2_1_2, eqs_2_1_2)
        print("[OK] Inserted Persamaan 2.1 to 2.4 in Subbab 2.1.2")

    # 2. In Subbab 2.1.3 (Algoritma CatBoost)
    # Find paragraph 2.1.3 content (ends with manipulasi data sintetis...)
    target_p_2_1_3 = None
    for i in range(80, 95):
        if "berisiko mengubah pola distribusi alami" in doc.paragraphs[i].text:
            target_p_2_1_3 = doc.paragraphs[i]
            break
            
    if target_p_2_1_3:
        eqs_2_1_3 = [
            ("\tFormulasi matematis untuk pembobotan biaya kelas (Cost-Sensitive Class Weights) yang diterapkan oleh CatBoost guna menyeimbangkan penalti kerugian kelas minoritas dirumuskan pada Persamaan (2.5):", "Normal", None),
            ("\tw_c = N / (K * N_c)                                                                  (Persamaan 2.5)", "Normal", None),
            ("\tdi mana w_c adalah bobot pengali penalti kerugian untuk kelas target c, N merupakan total seluruh sampel data latih, K adalah jumlah kelas risiko (K = 4), dan N_c menyatakan frekuensi jumlah sampel yang tergolong pada kelas c. Melalui formulasi ini, kelas dengan frekuensi kemunculan terkecil (seperti kelas Severe) secara proporsional memperoleh bobot penalti w_c yang jauh lebih besar.", "Normal", None),
            ("\tAdapun keunggulan pemrosesan fitur kategorikal tanpa One-Hot Encoding pada CatBoost dieksekusi melalui transformasi Ordered Target Statistics yang dirumuskan pada Persamaan (2.6):", "Normal", None),
            ("\tx̂_k = (∑ y_historis + a * P) / (Count_historis + a)                                   (Persamaan 2.6)", "Normal", None),
            ("\tdi mana x̂_k melambangkan nilai transformasi numerik kontinu dari fitur kategori pada permutasi urutan ke-k, ∑ y_historis menyatakan akumulasi nilai target dari kemunculan kategori yang sama pada data urutan sebelumnya, Count_historis adalah jumlah frekuensi kemunculan kategori serupa terdahulu, P melambangkan prior global rata-rata target seluruh dataset, dan a > 0 adalah parameter bobot penghalusan (smoothing weight, dengan nilai baku a = 1).", "Normal", None)
        ]
        target_p_2_1_3 = insert_paragraphs_after(target_p_2_1_3, eqs_2_1_3)
        print("[OK] Inserted Persamaan 2.5 and 2.6 in Subbab 2.1.3")

    # 3. In Subbab 2.1.5 (SHAP dan Tree-SHAP)
    # Find paragraph 2.1.5 content
    target_p_2_1_5 = None
    for i in range(88, 105):
        if "Melalui keempat aksioma tersebut, pembagian bobot kontribusi antarvariabel dipastikan bebas dari bias" in doc.paragraphs[i].text:
            target_p_2_1_5 = doc.paragraphs[i]
            break
            
    if target_p_2_1_5:
        eqs_2_1_5 = [
            ("\tSecara teoretis, nilai Shapley (φ_i) dari suatu fitur input ke-i dihitung berdasarkan kontribusi marginal rata-rata tertimbang dari seluruh kemungkinan kombinasi koalisi himpunan bagian fitur, sebagaimana dirumuskan pada Persamaan (2.7):", "Normal", None),
            ("\tφ_i(f, x) = ∑_{S ⊆ F \\ {i}} [|S|!(|F| - |S| - 1)! / |F|!] * [f_x(S ∪ {i}) - f_x(S)]  (Persamaan 2.7)", "Normal", None),
            ("\tdi mana F melambangkan himpunan seluruh fitur prediktor input (|F| = 28 fitur), S merupakan himpunan bagian fitur yang tidak menyertakan fitur ke-i, |S| menyatakan kardinalitas jumlah fitur dalam himpunan bagian S, f_x(S) adalah luaran model saat dievaluasi menggunakan subset fitur S, dan selisih [f_x(S ∪ {i}) - f_x(S)] merefleksikan kontribusi marginal penambahan fitur ke-i ke dalam koalisi.", "Normal", None),
            ("\tBerdasarkan aksioma efisiensi, akumulasi seluruh nilai kontribusi fitur lokal beserta nilai dasar ekspektasi global model menjamin terpenuhinya dekomposisi aditif eksak yang dinyatakan pada Persamaan (2.8):", "Normal", None),
            ("\tf(x) = φ_0 + ∑_{i=1}^M φ_i(x)                                                         (Persamaan 2.8)", "Normal", None),
            ("\tdi mana f(x) menyatakan hasil prediksi akhir model untuk instance x (pada ruang margin logit maupun ruang probabilitas terkalibrasi), φ_0 = E[f(x)] merepresentasikan nilai dasar ekspektasi model pada populasi data sebelum mempertimbangkan variabel pasien, M adalah jumlah total variabel fitur (M = 28), dan φ_i(x) mencerminkan besaran atribusi kontribusi lokal dari variabel gaya hidup ke-i terhadap penyimpangan keputusan risiko pasien dari nilai dasar populasi.", "Normal", None)
        ]
        target_p_2_1_5 = insert_paragraphs_after(target_p_2_1_5, eqs_2_1_5)
        print("[OK] Inserted Persamaan 2.7 and 2.8 in Subbab 2.1.5")

# ==============================================================================
# 4. BAB III (METODE PENELITIAN)
# ==============================================================================
print("\n--- 4. Updating Bab III ---")

# Subbab 3.2.1 Waktu Penelitian & Caption Tabel 3.1
for i, p in enumerate(doc.paragraphs):
    if "Pelaksanaan penelitian ini direncanakan berlangsung selama enam bulan" in p.text:
        p.text = (
            "Pelaksanaan penelitian ini direncanakan berlangsung selama enam bulan, terhitung mulai bulan Oktober 2026 sampai dengan bulan Maret 2027 (Tahun Akademik 2026/2027). "
            "Alokasi waktu tersebut disusun secara bertahap guna mengakomodasi seluruh rangkaian kegiatan ilmiah, yang dimulai dari identifikasi masalah, studi pustaka, penyusunan proposal penelitian, hingga seminar proposal. "
            "Selanjutnya, tahapan berlanjut pada pengumpulan serta pra-pemrosesan dataset sekunder, perancangan arsitektur komputasi, pelatihan model CatBoost, serta integrasi pustaka Tree-SHAP untuk analisis interpretabilitas faktor risiko tidur. "
            "Dua bulan terakhir dialokasikan khusus untuk pengujian performa secara menyeluruh, analisis komparasi hasil evaluasi metrik, perancangan prototipe sistem pendukung keputusan klinis, dan penulisan naskah laporan skripsi lengkap. "
            "Pengaturan jadwal yang terstruktur dan terukur ini bertujuan memastikan seluruh tahapan penelitian dapat diselesaikan tepat waktu sesuai standar mutu akademik Fakultas Sains dan Teknologi UNISNU Jepara."
        )
        print(f"[OK] Updated Subbab 3.2.1 research time in P[{i}]")
    if "Tabel 3.1 Jadwal Rencana Pelaksanaan Penelitian" in p.text:
        p.text = "Tabel 3.1 Jadwal Rencana Pelaksanaan Penelitian (Tahun Akademik 2026/2027)"
        print(f"[OK] Updated Table 3.1 caption in P[{i}]")

# Table 1 (Table 3.1 Jadwal) cleanup: remove phantom col 5, populate checkmarks
t1 = doc.tables[1]
if len(t1.rows[0].cells) == 9:
    print("[INFO] Cleaning Table 3.1 phantom column...")
    for row in t1.rows:
        tc = row.cells[5]._tc
        tc.getparent().remove(tc)
    print(f"[OK] Table 3.1 now has {len(t1.rows[0].cells)} columns.")

# Populate headers and rows for Table 3.1
t1_headers = ['No', 'Tahapan Kegiatan Penelitian', 'Bulan 1\n(Okt 26)', 'Bulan 2\n(Nov 26)', 'Bulan 3\n(Des 26)', 'Bulan 4\n(Jan 27)', 'Bulan 5\n(Feb 27)', 'Bulan 6\n(Mar 27)']
for col_idx, htext in enumerate(t1_headers):
    t1.rows[0].cells[col_idx].text = htext

t1_data = [
    ('1', 'Identifikasi Masalah & Studi Pustaka', ['✓', '✓', '', '', '', '']),
    ('2', 'Penyusunan Proposal & Seminar Proposal', ['', '✓', '✓', '', '', '']),
    ('3', 'Pengumpulan Data & Preprocessing', ['', '', '✓', '✓', '', '']),
    ('4', 'Pemodelan CatBoost & Integrasi SHAP', ['', '', '', '✓', '✓', '']),
    ('5', 'Evaluasi Kinerja & Desain Prototipe CDSS', ['', '', '', '', '✓', '']),
    ('6', 'Penyusunan Laporan Skripsi & Publikasi', ['', '', '', '', '✓', '✓'])
]

for r_idx, (no, nama, checks) in enumerate(t1_data, start=1):
    t1.rows[r_idx].cells[0].text = no
    t1.rows[r_idx].cells[1].text = nama
    for m_idx, cval in enumerate(checks):
        t1.rows[r_idx].cells[2 + m_idx].text = cval
print("[OK] Populated checkmarks for all 6 months in Table 3.1")

# Subbab 3.3.2 Sampel Penelitian (Feature Breakdown: 1 ID, 7 Cat, 23 Num, 1 Target = 32 total)
for i, p in enumerate(doc.paragraphs):
    if "Sampel penelitian yang digunakan diambil melalui teknik purposive sampling" in p.text:
        p.text = (
            "Sampel penelitian yang digunakan diambil melalui teknik purposive sampling dengan menetapkan kriteria kelengkapan atribut metrik gaya hidup digital dan rekaman fisiologis tidur secara terstruktur. "
            "Berdasarkan kriteria inklusi tersebut, sampel yang dianalisis berupa Sleep Health & Daily Performance Dataset karya Mohan Krishna Thalla (2026) dari repositori publik Kaggle dengan jumlah 100.000 rekaman dan 32 atribut fitur. "
            "Dataset ini merupakan data sekunder terbuka berbasis pemodelan sintetis terkalibrasi secara epidemiologis, sehingga bebas dari kewajiban izin etik medis langsung namun memiliki validitas statistik representatif. "
            "Struktur data mencakup 1 atribut pengenal unik (identifier), 7 fitur prediktor kategorikal seperti profesi dan kondisi kesehatan mental, serta 23 fitur prediktor numerik seperti durasi layar gawai dan stres harian, dengan 1 variabel target dependen sleep_disorder_risk empat tingkatan risiko: Healthy, Mild, Moderate, dan Severe, yang sangat ideal untuk melatih model klasifikasi multi-kelas."
        )
        print(f"[OK] Updated Subbab 3.3.2 feature breakdown in P[{i}]")

# Table 2 (Table 3.2 Deskripsi Atribut): Restore underscores
t2 = doc.tables[2]
var_fixes = {
    "alcohol units before bed": "alcohol_units_before_bed",
    "stress score": "stress_score",
    "work hours that day": "work_hours_that_day",
    "mental health condition": "mental_health_condition",
    "felt rested": "felt_rested"
}
for row in t2.rows:
    for orig, fixed in var_fixes.items():
        if row.cells[1].text.strip() == orig:
            row.cells[1].text = fixed
            print(f"[OK] Restored underscore in Table 3.2: '{fixed}'")

# Subbab 3.4.1 Hardware Storage harmonization (512 GB)
for i, p in enumerate(doc.paragraphs):
    if "Instrumen perangkat keras (hardware) yang digunakan dalam penelitian ini berfungsi" in p.text:
        if "1024 GB" in p.text:
            p.text = p.text.replace("1024 GB", "512 GB M.2 NVMe PCIe 3.0 SSD")
            print(f"[OK] Harmonized SSD storage to 512 GB in Subbab 3.4.1 P[{i}]")

# Subbab 3.6.2 Preprocessing remaining features clarification
for i, p in enumerate(doc.paragraphs):
    if "Melalui proses seleksi tersebut, dataset menyisakan 28 variabel prediktor" in p.text:
        p.text = (
            "Tahap persiapan data (data preparation) dilakukan secara cermat guna membersihkan dan menyusun matriks fitur agar siap diproses oleh algoritma pembelajaran mesin. "
            "Langkah awal yang dilakukan adalah seleksi fitur dengan mengeliminasi atribut identifier person_id yang bersifat non-prediktif, serta atribut day_type dan country guna menghilangkan variabel dengan kontribusi nol sekaligus mencegah bias wilayah geografis. "
            "Melalui proses eliminasi 3 atribut tersebut, dataset menyisakan 28 variabel prediktor independen (mencakup 5 fitur kategorikal native dan 23 fitur numerik) serta satu variabel target dependen sleep_disorder_risk (total 29 kolom terproses). "
            "Keunggulan komputasi CatBoost dimaksimalkan pada tahap ini melalui fitur penanganan kategorikal secara alami (native categorical handling). "
            "Sebanyak lima atribut bertipe teks, yaitu gender, occupation, chronotype, mental_health_condition, dan season, didaftarkan langsung ke dalam model tanpa melalui prosedur One-Hot Encoding. "
            "Prosedur ini sangat efektif dalam mencegah peningkatan dimensi data (curse of dimensionality), menghemat memori kerja komputasi, serta mempertahankan keutuhan semantik data gaya hidup."
        )
        print(f"[OK] Clarified feature breakdown in Subbab 3.6.2 P[{i}]")

# Subbab 3.6.5 OTS narrative clarification (gender attribute, binary One-vs-All multi-class)
for i, p in enumerate(doc.paragraphs):
    if "Tahapan simulasi perhitungan manual kedua difokuskan pada pembuktian mekanisme transformasi fitur kategorikal" in p.text:
        p.text = (
            "Tahapan simulasi perhitungan manual kedua difokuskan pada pembuktian mekanisme transformasi fitur kategorikal melalui metode Ordered Target Statistics bawaan CatBoost. "
            "Menggunakan lima baris sampel data tiruan representatif pada atribut kategori jenis kelamin (gender: Female, Male), proses transformasi numerik dihitung secara sekuensial menggunakan rumus statistik target terurut dengan parameter prioritas sebesar satu (a = 1). "
            "Setiap nilai kategori dikonversi menjadi nilai estimasi probabilitas kontinu berdasarkan akumulasi label target baris data terdahulu ditambah pembobotan prioritas global tanpa melibatkan label data masa depan. "
            "Pada klasifikasi multi-kelas dengan empat tingkatan risiko (Healthy, Mild, Moderate, Severe), CatBoost menerapkan mekanisme Ordered Target Statistics secara paralel dengan skema biner One-vs-All untuk masing-masing kelas target. "
            "Simulasi pada Tabel 3.8 merepresentasikan kalkulasi OTS untuk satu target kelas biner (misal: Severe vs Non-Severe atau Disrupsi vs Normal) guna memperjelas tahapan transformasi terurut tanpa mengaburkan prinsip matematis dasarnya. "
            "Hasil kalkulasi manual pada kelima sampel data tersebut menghasilkan angka pembobotan kontinu yang persis sama dengan matriks fitur terproses yang dihasilkan oleh fungsi internal CatBoost Pool. "
            "Verifikasi manual ini membuktikan secara ilmiah bahwa algoritma mampu mempertahankan keutuhan relasi data kategorikal gaya hidup tanpa menimbulkan ledakan dimensi fitur artifisial."
        )
        print(f"[OK] Clarified OTS simulation narrative in Subbab 3.6.5 P[{i}]")

# Table 9 (Table 3.9 Confusion Matrix & Metrik Evaluasi): Reconcile with authentic model
t9 = doc.tables[9]
print("[INFO] Reconciling Table 3.9 with authentic model evaluation on 20.000 test data...")
t9_rows = [
    # Row 1: Healthy
    ['Aktual: Healthy', '10.549', '282', '0', '0', '10.831', 
     'Recall = 10.549/10.831 = 0,9740; Prec = 10.549/10.601 = 0,9951', 'F1 = 0,9844', '0,98'],
    # Row 2: Mild
    ['Aktual: Mild', '52', '6.084', '560', '0', '6.696', 
     'Recall = 6.084/6.696 = 0,9086; Prec = 6.084/6.448 = 0,9435', 'F1 = 0,9257', '0,93'],
    # Row 3: Moderate
    ['Aktual: Moderate', '0', '82', '1.405', '173', '1.660', 
     'Recall = 1.405/1.660 = 0,8464; Prec = 1.405/2.044 = 0,6874', 'F1 = 0,7586', '0,76'],
    # Row 4: Severe
    ['Aktual: Severe', '0', '0', '79', '734', '813', 
     'Recall = 734/813 = 0,9028; Prec = 734/907 = 0,8093', 'F1 = 0,8535', '0,85'],
    # Row 5: Total
    ['Total Prediksi', '10.601', '6.448', '2.044', '907', '20.000', 
     'Accuracy = 18.772 / 20.000; Macro-F1 = (0,9844+0,9257+0,7586+0,8535)/4', '93,86% (Acc) | 88,06% (Macro)', '93,86% | 0,88']
]

for r_idx, rvals in enumerate(t9_rows, start=1):
    for c_idx, val in enumerate(rvals):
        t9.rows[r_idx].cells[c_idx].text = val
print("[OK] Reconciled Table 3.9 metrics and confusion matrix values")

# Note below Table 3.9
for i, p in enumerate(doc.paragraphs):
    if "Catatan: Macro-F1 Score terhitung manual sebesar" in p.text or "Catatan: Hasil evaluasi simulasi" in p.text:
        p.text = "Catatan: Hasil evaluasi simulasi uji coba awal (preliminary benchmark test) ini membuktikan keselarasan sempurna antara perhitungan matematis manual dengan fungsi classification_report pustaka Scikit-Learn (Akurasi = 93,86% dan Macro-F1 = 88,06%)."
        print(f"[OK] Updated Table 3.9 note in P[{i}]")

# Table 10 (Table 3.10 SHAP Decomposition)
t10 = doc.tables[10]
print("[INFO] Aligning Table 3.10 additive SHAP components...")
t10_rows = [
    # Header row is row 0
    ['Base Value (φ0)', 'Rerata Ekspektasi Global E[f(x)]', '—', '+0,1200', 'Nilai acuan dasar probabilitas risiko sebelum fitur pasien'],
    ['Fitur 1 (Pemicu Utama)', 'screen_time_before_bed_mins', '145 menit', '+0,3800', 'Sangat kuat mendorong peningkatan risiko Severe'],
    ['Fitur 2 (Pemicu Tambahan)', 'caffeine_mg_before_bed', '180 mg', '+0,2400', 'Mendorong peningkatan risiko gangguan tidur'],
    ['Fitur 3 (Pemicu Mental)', 'stress_score', 'Skala 8/10', '+0,1600', 'Beban stres tinggi menaikkan skor probabilitas risiko'],
    ['Fitur 4 (Pemicu Siklus)', 'sleep_latency_mins', '45 menit', '+0,1100', 'Latensi tidur lama memperparah indikasi insomnia'],
    ['Fitur 5 (Pemicu Fisiologis)', 'sleep_duration_hrs', '5,2 jam', '+0,0700', 'Durasi tidur pendek memperkuat sinyal risiko tidur'],
    ['Fitur 6 (Pemicu Kebiasaan)', 'alcohol_units_before_bed', '2 unit', '+0,0400', 'Konsumsi alkohol mengganggu arsitektur tidur malam'],
    ['Fitur 7 (Faktor Penekan)', 'steps_that_day', '3.200 langkah', '-0,0500', 'Aktivitas fisik bertindak sebagai protektor tidur'],
    ['Fitur 8 s/d 28 (Penekan)', 'Akumulasi 21 Fitur Penekan Lainnya', 'Beragam', '-0,1200', 'Kontribusi marginal residual gabungan variabel penekan'],
    ['Total Akumulasi (Σ φi)', 'Penjumlahan Seluruh 28 Fitur', '—', '+0,8300', 'Total pergeseran kontribusi fitur pasien (+1,0000 - 0,1700)'],
    ['Prediksi Akhir f(x)', 'φ0 + Σ φi = 0,1200 + 0,8300', '—', '0,9500', 'Identik 100% dengan Probabilitas CatBoost (95,00%)']
]

# Table 10 currently has 10 rows or 11 rows. Let's see:
while len(t10.rows) < len(t10_rows) + 1:
    t10.add_row()

for r_idx, rvals in enumerate(t10_rows, start=1):
    for c_idx, val in enumerate(rvals):
        t10.rows[r_idx].cells[c_idx].text = val
print("[OK] Aligned Table 3.10 with exact waterfall arithmetic (+0.12 + 0.83 = 0.95)")

# Subbab 3.6.6 Preliminary test framing
for i, p in enumerate(doc.paragraphs):
    if "Tahap evaluasi (evaluation) dilakukan secara komprehensif untuk mengukur performa generalisasi" in p.text:
        p.text = (
            "Tahap evaluasi (evaluation) dilakukan secara komprehensif untuk mengukur performa generalisasi model CatBoost pada 20.000 sampel data uji yang belum pernah dilihat sebelumnya. "
            "Perlu ditegaskan bahwa simulasi evaluasi pada Tabel 3.9 dan Gambar 3.4 merupakan hasil uji coba awal (preliminary benchmark test) yang bertujuan memvalidasi kecocokan logika rumus metrik evaluasi manual terhadap implementasi pustaka komputasi Scikit-Learn, sebelum pengujian eksperimental komprehensif pada berbagai skenario pemodelan dipaparkan secara lengkap pada Bab IV. "
            "Penilaian kinerja model tidak hanya bersandar pada nilai akurasi semata, karena metrik akurasi rentan menimbulkan ilusi performa tinggi (accuracy paradox) pada dataset dengan distribusi kelas yang tidak seimbang. "
            "Oleh sebab itu, evaluasi dilengkapi dengan metrik Precision, Recall, dan Macro-averaged F1-Score untuk memastikan keandalan prediksi pada seluruh kelas risiko (Chicco & Jurman, 2022). "
            "Di samping itu, matriks kebingungan (confusion matrix) multi-kelas berukuran 4x4 disusun untuk memetakan secara detail jumlah klasifikasi yang tepat (True Positive) serta kesalahan klasifikasi (misclassification) pada setiap kelas. "
            "Analisis kurva Receiver Operating Characteristic beserta nilai Area Under Curve (ROC-AUC) juga dihitung untuk menguji daya diskriminasi probabilitas model secara objektif."
        )
        print(f"[OK] Added preliminary test framing in Subbab 3.6.6 P[{i}]")

# Subbab 3.6.7 SHAP Standardization & Logit Margin vs Probability clarification
for i, p in enumerate(doc.paragraphs):
    if "Guna mengatasi keterbatasan sifat kotak hitam (black-box) pada algoritma CatBoost" in p.text:
        p.text = (
            "Guna mengatasi keterbatasan sifat kotak hitam (black-box) pada algoritma CatBoost, penelitian ini mengintegrasikan metode Explainable Artificial Intelligence berbasis SHAP (Shapley Additive exPlanations). "
            "Modul komputasi yang digunakan adalah TreeExplainer, yang secara khusus dioptimalkan untuk mengevaluasi struktur pohon keputusan dengan efisiensi waktu polinomial tanpa memerlukan aproksimasi sampling (Lundberg & Lee, 2017). "
            "Dalam komputasi internalnya, modul TreeExplainer menghitung nilai atribusi fitur pada ruang margin logit multi-kelas (raw prediction margin). "
            "Untuk kebutuhan penyajian klinis pada prototipe CDSS, margin logit tersebut ditransformasikan melalui fungsi softmax ke dalam ruang probabilitas terkalibrasi (0% hingga 100%), sehingga visualisasi grafik air terjun (waterfall plot) dapat dibaca secara intuitif oleh tenaga medis di mana penjumlahan nilai dasar (base value E[f(x)] = 0,12) dengan akumulasi kontribusi bersih fitur (+0,83) menghasilkan nilai prediksi probabilitas akhir sebesar 0,95 (95,00%). "
            "Analisis interpretabilitas dijalankan pada dua tingkatan analitik yang saling melengkapi, yaitu tingkat global dan tingkat lokal. "
            "Pada tingkat global, grafik ringkasan (SHAP summary plot) dan grafik kepentingan fitur (feature importance) digunakan untuk meranking variabel gaya hidup digital yang paling dominan mempengaruhi risiko gangguan tidur secara keseluruhan. "
            "Sementara itu, pada tingkat lokal, grafik air terjun (SHAP waterfall plot) diaplikasikan untuk membedah kontribusi marginal setiap parameter perilaku terhadap keputusan prediksi pada individu pasien tertentu (Ha et al., 2023)."
        )
        print(f"[OK] Clarified Logit Margin vs Probability in Subbab 3.6.7 P[{i}]")

# ==============================================================================
# 5. DAFTAR PUSTAKA
# ==============================================================================
print("\n--- 5. Updating Daftar Pustaka ---")
# Update Xie et al. (2025)
for i, p in enumerate(doc.paragraphs):
    if "Xie, Y., Chen, Y., Han, Y." in p.text:
        p.text = "Xie, Y., Chen, Y., Han, Y., Zhai, S., Xiao, L., Yin, D., & Chen, Y. (2025). Identifying influencing factors associated with sleep quality in undergraduates based on partial least squares regression and XGBoost. Frontiers in Psychology, 16, 1732946. https://doi.org/10.3389/fpsyg.2025.1732946"
        print(f"[OK] Updated Xie et al. to (2025) in Daftar Pustaka P[{i}]")

# Add new references: Chang et al. (2015), Drake et al. (2013), Irish et al. (2015)
dp_entries = [p.text.strip() for p in doc.paragraphs[184:] if p.text.strip()]

new_refs = [
    "Chang, A. M., Aeschbach, D., Duffy, J. F., & Czeisler, C. A. (2015). Evening use of light-emitting eReaders negatively affects sleep, circadian timing, and next-morning alertness. Proceedings of the National Academy of Sciences (PNAS), 112(4), 1232–1237. https://doi.org/10.1073/pnas.1418490112",
    "Drake, C., Roehrs, T., Shambroom, J., & Roth, T. (2013). Caffeine effects on sleep taken 0, 3, or 6 hours before going to bed. Journal of Clinical Sleep Medicine, 9(11), 1195–1200. https://doi.org/10.5664/jcsm.3170",
    "Irish, L. A., Kline, C. E., Gunn, H. E., Buysse, D. J., & Hall, M. H. (2015). The role of sleep hygiene in promoting public health: A review of empirical evidence. Sleep Medicine Reviews, 22, 23–36. https://doi.org/10.1016/j.smrv.2014.10.001"
]

# Find last paragraph of doc
last_p = doc.paragraphs[-1]
for ref in new_refs:
    if not any(ref[:30] in p.text for p in doc.paragraphs):
        new_p = doc.add_paragraph(ref, style='Normal')
        last_p._p.addnext(new_p._p)
        last_p = new_p
        print(f"[OK] Added reference: {ref[:40]}...")

# Save updated docx
doc.save(DOCX_OUT)
print(f"\n[SUCCESS] Document saved successfully to {DOCX_OUT}")

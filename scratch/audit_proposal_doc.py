import docx
import re
import os

doc = docx.Document("temp_inspect_proposal.docx")

print("=== AUDIT: PROPOSAL BAB 1-3.docx ===")
print("Total Paragraphs:", len(doc.paragraphs))
print("Total Tables:", len(doc.tables))

# Banned words check
banned_words = [
    r"\bsaya\b", r"\bkami\b", r"\bpenulis\b", r"\bkita\b",
    r"\bsangat canggih\b", r"\bluar biasa\b", r"\btebakan\b",
    r"\bbisa dibilang\b", r"\balat bernama\b", r"\bcanggih\b"
]

print("\n--- 1. BANNED WORDS AUDIT ---")
banned_found = []
for i, p in enumerate(doc.paragraphs):
    for bw in banned_words:
        matches = re.findall(bw, p.text, re.IGNORECASE)
        if matches:
            banned_found.append((i, p.text.strip()[:80], matches))

print(f"Banned words occurrences: {len(banned_found)}")
for i, snippet, m in banned_found[:10]:
    print(f"  P[{i}]: matches={m} | {snippet}...")

# Check Chen vs Xie
print("\n--- 2. CHEN VS XIE AUDIT ---")
chen_occurrences = []
xie_occurrences = []
for i, p in enumerate(doc.paragraphs):
    if "Chen et al." in p.text or "Chen, T." in p.text or "Chen et al.," in p.text:
        chen_occurrences.append((i, p.text.strip()[:90]))
    if "Xie et al." in p.text or "Xie, Y." in p.text or "Xie et al.," in p.text:
        xie_occurrences.append((i, p.text.strip()[:90]))

print(f"Chen occurrences in paragraphs: {len(chen_occurrences)}")
for i, snippet in chen_occurrences:
    print(f"  P[{i}]: {snippet}")

print(f"Xie occurrences in paragraphs: {len(xie_occurrences)}")
for i, snippet in xie_occurrences:
    print(f"  P[{i}]: {snippet}")

# Check Table 0 (Table 2.1)
print("\n--- 3. TABLE 2.1 (TABLE 0) AUDIT ---")
t0 = doc.tables[0]
for r_idx, row in enumerate(t0.rows):
    vals = [c.text.strip().replace("\n", " ") for c in row.cells]
    print(f"  Row {r_idx}: {vals[0]} | {vals[1]} | {vals[2][:50]}...")

# Check Bab 1 Latar Belakang Paragraph Word Counts
print("\n--- 4. BAB 1 WORD COUNTS (P[7] to P[12]) ---")
for i in range(7, 13):
    words = len(doc.paragraphs[i].text.split())
    print(f"  P[{i}]: {words} words | {doc.paragraphs[i].text.strip()[:60]}...")

# Check Bab 2 Paragraph Word Counts
print("\n--- 5. BAB 2 PARAGRAPH WORD COUNTS ---")
for i in range(31, 62):
    p = doc.paragraphs[i]
    words = len(p.text.split())
    if words > 10:
        print(f"  P[{i}] ({p.style.name}): {words} words | {p.text.strip()[:60]}...")

# Check DAFTAR PUSTAKA
print("\n--- 6. DAFTAR PUSTAKA AUDIT ---")
start_dp = -1
for i, p in enumerate(doc.paragraphs):
    if "DAFTAR PUSTAKA" in p.text.upper():
        start_dp = i
        break

if start_dp != -1:
    dp_entries = [p.text.strip() for p in doc.paragraphs[start_dp + 1:] if p.text.strip()]
    print(f"Total DAFTAR PUSTAKA entries: {len(dp_entries)}")
    recent = sum(1 for e in dp_entries if any(y in e for y in ["(2021)", "(2022)", "(2023)", "(2024)", "(2025)", "(2026)"]))
    print(f"References from 2021-2026: {recent} / {len(dp_entries)} ({recent/len(dp_entries)*100:.2f}%)")
    
    # Check if Chen is in DP
    chen_dp = [e for e in dp_entries if "Chen, T." in e or "1373504" in e]
    print(f"Erroneous Chen entry in DP: {len(chen_dp)}")
    if chen_dp:
        print(f"  {chen_dp[0]}")
        
    # Check if Xie is in DP
    xie_dp = [e for e in dp_entries if e.startswith("Xie, Y.")]
    print(f"Correct Xie entry in DP: {len(xie_dp)}")
    if xie_dp:
        print(f"  {xie_dp[0]}")

import docx
from docx.shared import Inches, Cm, Pt

doc = docx.Document("temp_inspect_proposal.docx")

print("=== MARGINS AND PAGE SETUP ===")
for i, sec in enumerate(doc.sections):
    print(f"Section {i}:")
    print(f"  Top: {sec.top_margin.cm:.2f} cm")
    print(f"  Bottom: {sec.bottom_margin.cm:.2f} cm")
    print(f"  Left: {sec.left_margin.cm:.2f} cm")
    print(f"  Right: {sec.right_margin.cm:.2f} cm")
    print(f"  Page Width: {sec.page_width.cm:.2f} cm, Height: {sec.page_height.cm:.2f} cm")

print("\n=== STYLES AND RUNS SAMPLING ===")
# Sample 10 paragraphs across Bab 1, 2, 3
sample_indices = [7, 14, 34, 40, 52, 65, 92, 107, 130, 140]
for idx in sample_indices:
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        run_fonts = set(r.font.name for r in p.runs if r.font.name)
        run_sizes = set(r.font.size.pt for r in p.runs if r.font.size)
        print(f"P[{idx:3d}] ({p.style.name:15s}): Font={run_fonts or 'default'}, Size={run_sizes or 'default'}, Spacing={p.paragraph_format.line_spacing}, Indent={p.paragraph_format.first_line_indent.inches if p.paragraph_format.first_line_indent else 'none'}")

import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('scratch/temp_verify_final.docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

print("\n--- SECTION HEADINGS AUDIT ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if any(txt.startswith(prefix) for prefix in ["BAB ", "1.", "2.", "3.", "DAFTAR "]):
        print(f"P[{i:3d}]: {txt}")

print("\n--- TABLE 2.1 (STATE OF THE ART) AUDIT ---")
t0 = doc.tables[0]
for r_idx, r in enumerate(t0.rows):
    cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
    print(f"  Row {r_idx}: {cells[1] if len(cells) > 1 else ''} | {cells[2][:40] if len(cells) > 2 else ''}")

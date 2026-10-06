import docx
import os

target = "BAB_1_1438_temp.docx"
print(f"File exists: {os.path.exists(target)}, size: {os.path.getsize(target)}")

with open(target, "rb") as f:
    doc = docx.Document(f)
print("Paragraphs count:", len(doc.paragraphs))
print("Tables count:", len(doc.tables))
print("Sections count:", len(doc.sections))

print("\n--- OUTLINE / HEADINGS DETECTED ---")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    # Check if heading or starts with BAB or uppercase title
    if p.style.name.startswith("Heading") or text.upper().startswith("BAB") or any(text.upper().startswith(x) for x in ["KATA PENGANTAR", "DAFTAR ISI", "DAFTAR TABEL", "DAFTAR GAMBAR", "DAFTAR PUSTAKA"]):
        print(f"P[{i}] (Style: {p.style.name}): {text[:100]}")
    elif len(text) < 60 and (text.startswith("1.") or text.startswith("2.") or text.startswith("3.") or text.startswith("A.") or text.startswith("B.") or text.startswith("C.")):
        print(f"P[{i}] (Style: {p.style.name}): {text[:100]}")

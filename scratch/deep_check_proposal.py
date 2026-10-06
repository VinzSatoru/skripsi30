import docx
import re

doc = docx.Document("PROPOSAL BAB 1-3.docx")

print("=== 1. SCAN FOR NON-BREAKING SPACES OR SPECIAL CHARS ===")
nbsp_count = 0
for i, p in enumerate(doc.paragraphs):
    if "\xa0" in p.text:
        nbsp_count += p.text.count("\xa0")

print(f"Total \\xa0 found in paragraphs: {nbsp_count}")

# Check tables as well
nbsp_tables = 0
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if "\xa0" in cell.text:
                nbsp_tables += cell.text.count("\xa0")
print(f"Total \\xa0 found in tables: {nbsp_tables}")

print("\n=== 2. SCAN FOR ANY REMAINING 'Chen' IN DOC ===")
for i, p in enumerate(doc.paragraphs):
    if "Chen" in p.text:
        print(f"P[{i}]: {p.text[:90]}")
for t_idx, t in enumerate(doc.tables):
    for r_idx, row in enumerate(t.rows):
        for c_idx, cell in enumerate(row.cells):
            if "Chen" in cell.text:
                print(f"Table {t_idx} [{r_idx},{c_idx}]: {cell.text.strip()[:60]}")

print("\n=== 3. CHECK CITATION CONSISTENCY IN BAB 2 ===")
for i in range(31, 62):
    p = doc.paragraphs[i]
    citations = re.findall(r'\([A-Za-z\s&.,]+?\d{4}\)', p.text)
    if citations:
        print(f"P[{i}]: {citations}")

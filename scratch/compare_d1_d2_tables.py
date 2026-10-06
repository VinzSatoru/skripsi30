import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

d1 = docx.Document('PROPOSAL BAB 1-3.docx')
d2 = docx.Document('scratch/temp_read.docx')

print("D1 (PROPOSAL BAB 1-3.docx) tables:")
for i, t in enumerate(d1.tables):
    h = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells[:3]]
    print(f"  T[{i}]: {len(t.rows)}x{len(t.columns)} | {h}")

print("\nD2 (PROPOSAL_BAB_1-3_Revisi.docx) tables:")
for i, t in enumerate(d2.tables):
    h = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells[:3]]
    print(f"  T[{i}]: {len(t.rows)}x{len(t.columns)} | {h}")

import sys
import os
import shutil
import docx
from docx.shared import Pt, Inches

sys.stdout.reconfigure(encoding='utf-8')

base_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(base_dir)

src_file = os.path.join(base_dir, "temp_read.docx")
final_file = os.path.join(parent_dir, "PROPOSAL_BAB_1-3_FINAL_SEMPURNA.docx")

with open(src_file, "rb") as f:
    doc = docx.Document(f)

# 1. Clean trailing tab in P[1]
if len(doc.paragraphs) > 1 and "\t" in doc.paragraphs[1].text:
    doc.paragraphs[1].text = doc.paragraphs[1].text.strip()
    print("[OK] P[1] cleaned trailing tab.")

# 2. Standardize Bab 3 Headings
heading_replacements = {
    "3.1. Jenis dan Pendekatan Penelitian": "3.1 Jenis dan Pendekatan Penelitian",
    "3.1.1. Jenis Penelitian": "3.1.1 Jenis Penelitian",
    "3.1.2. Pendekatan Penelitian": "3.1.2 Pendekatan Penelitian",
    "3.2. Waktu dan Tempat Penelitian": "3.2 Waktu dan Tempat Penelitian",
    "3.2.1. Waktu Penelitian": "3.2.1 Waktu Penelitian",
    "3.2.2. Tempat Penelitian": "3.2.2 Tempat Penelitian",
    "3.3. Populasi dan Sampel Penelitian": "3.3 Populasi dan Sampel Penelitian",
    "3.3.1. Populasi Penelitian": "3.3.1 Populasi Penelitian",
    "3.3.2. Sampel Penelitian": "3.3.2 Sampel Penelitian",
    "3.4. Instrumen Penelitian": "3.4 Instrumen Penelitian",
    "3.4.1. Perangkat Keras (Hardware)": "3.4.1 Perangkat Keras (Hardware)",
    "3.4.2. Perangkat Lunak (Software)": "3.4.2 Perangkat Lunak (Software)",
    "3.5. Teknik Pengumpulan Data": "3.5 Teknik Pengumpulan Data",
    "3.6. Teknik Analisis Data": "3.6 Teknik Analisis Data",
    "3.6.1. Pemahaman Bisnis dan Data (Business and Data Understanding)": "3.6.1 Pemahaman Bisnis dan Data (Business and Data Understanding)",
    "3.6.2. Pra-pemrosesan Data (Data Preprocessing)": "3.6.2 Pra-pemrosesan Data (Data Preprocessing)",
    "3.6.3. Pembagian Data dan Penanganan Ketidakseimbangan Kelas": "3.6.3 Pembagian Data dan Penanganan Ketidakseimbangan Kelas",
    "3.6.4. Pemodelan Machine Learning (CatBoost Classifier)": "3.6.4 Pemodelan Machine Learning (CatBoost Classifier)",
    "3.6.5. Simulasi Perhitungan Matematis Manual (Toy Example Workflow)": "3.6.5 Simulasi Perhitungan Matematis Manual (Toy Example Workflow)",
    "3.6.6. Evaluasi Kinerja Model": "3.6.6 Evaluasi Kinerja Model",
    "3.6.6. Interpretasi Model Berbasis XAI-SHAP": "3.6.7 Interpretasi Model Berbasis XAI-SHAP",
    "3.6.7. Perancangan Prototipe Sistem Pendukung Keputusan Klinis (Deployment)": "3.6.8 Perancangan Prototipe Sistem Pendukung Keputusan Klinis (Deployment)"
}

for p in doc.paragraphs:
    txt = p.text.strip()
    if txt in heading_replacements:
        new_txt = heading_replacements[txt]
        p.text = new_txt
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.bold = True
        print(f"[OK] Heading updated: '{txt}' -> '{new_txt}'")

doc.save(final_file)
print(f"[SUCCESS] Saved master document to {final_file}!")

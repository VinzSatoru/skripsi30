import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('PROPOSAL BAB 1-3.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
dp_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "DAFTAR PUSTAKA" in p.text.upper():
        dp_idx = i
        print(f"P[{i}]: {p.text}")
        break

if dp_idx != -1:
    print(f"DP entries count: {len(doc.paragraphs) - dp_idx - 1}")
    for j in range(dp_idx + 1, min(len(doc.paragraphs), dp_idx + 10)):
        print(f"  DP[{j}]: {doc.paragraphs[j].text[:80]}")

import docx
import shutil

shutil.copy2("PROPOSAL BAB 1-3.docx", "temp_inspect_proposal.docx")
with open("temp_inspect_proposal.docx", "rb") as f:
    doc = docx.Document(f)

# P[34] candidate:
p34_orig = doc.paragraphs[34].text.strip()
p34_mod = p34_orig.replace("psikologis manusia.", "psikologis manusia (Medic et al., 2017).")
w34 = len(p34_mod.split())
print(f"P[34] modified: {w34} words (Target: 120-130)")

# P[36] candidate:
p36_orig = doc.paragraphs[36].text.strip()
p36_mod = p36_orig.replace(
    "menciptakan kondisi kewaspadaan mental berlebihan yang menghambat ketenangan pikiran sebelum terlelap (Hapsari et al., 2024).",
    "menciptakan kondisi status kewaspadaan mental fisiologis berlebihan yang menghambat relaksasi pikiran sebelum terlelap (Widayati, 2024)."
)
w36 = len(p36_mod.split())
print(f"P[36] modified: {w36} words (Target: 120-130)")

# P[39] candidate:
p39_orig = doc.paragraphs[39].text.strip()
# Original has Kaya, 2025 and Taher & Ayon, 2024.
# Let's replace Kaya, 2025 with Mawardi et al., 2025, or add Mawardi et al., 2025:
p39_mod = p39_orig.replace(
    "(Kaya, 2025). Salah satu",
    "(Mawardi et al., 2025). Salah satu"
)
w39 = len(p39_mod.split())
print(f"P[39] modified: {w39} words (Target: 120-130)")

# P[49] candidate:
p49_orig = doc.paragraphs[49].text.strip()
p49_mod = p49_orig.replace(
    "(Zhang et al., 2025).",
    "(Zhang et al., 2025; Bhattarai et al., 2024)."
)
w49 = len(p49_mod.split())
print(f"P[49] modified: {w49} words (Target: 120-130)")

# P[67] candidate:
p67_orig = doc.paragraphs[67].text.strip()
p67_mod = p67_orig.replace(
    "(Martinez-Plumed et al., 2022; Lamaakal et al., 2025).",
    "(Martínez-Plumed et al., 2021; Schröer et al., 2021)."
)
w67 = len(p67_mod.split())
print(f"P[67] modified: {w67} words (Target: 120-130)")

# P[101] candidate:
p101_orig = doc.paragraphs[101].text.strip()
p101_mod = p101_orig.replace(
    "(Chakik et al., 2026).",
    "(El Chakik et al., 2026)."
)
w101 = len(p101_mod.split())
print(f"P[101] modified: {w101} words (Target: 120-130)")

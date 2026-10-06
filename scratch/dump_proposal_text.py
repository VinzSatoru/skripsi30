import sys
import os
import docx

sys.stdout.reconfigure(encoding='utf-8')

base_dir = os.path.dirname(os.path.abspath(__file__))
target_file = os.path.join(base_dir, "temp_read.docx")

with open(target_file, "rb") as f:
    doc = docx.Document(f)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

out_path = os.path.join(base_dir, "dump_proposal_text.txt")
with open(out_path, "w", encoding="utf-8") as out:
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            out.write(f"P[{i}]: {txt}\n\n")

print(f"Dumped text to {out_path}")

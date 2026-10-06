import sys
import os
import shutil
import docx

sys.stdout.reconfigure(encoding='utf-8')

new_title = "Judul: Penerapan Algoritma CatBoost dan SHAP untuk Klasifikasi Risiko Gangguan Tidur Berbasis Gaya Hidup"

files = [
    "PROPOSAL_BAB_1-3_FINAL_SEMPURNA.docx",
    "PROPOSAL BAB 1-3.docx",
    "PROPOSAL_BAB_1-3_Revisi.docx"
]

for fname in files:
    if os.path.exists(fname):
        with open(fname, "rb") as f:
            doc = docx.Document(f)
        
        # P[2] is the Judul paragraph
        if len(doc.paragraphs) > 2 and doc.paragraphs[2].text.startswith("Judul:"):
            doc.paragraphs[2].text = new_title
            print(f"[OK] Updated title in {fname}: {new_title}")
        
        doc.save(fname)

print("[SUCCESS] All proposal documents updated with new title (13 kata, tanpa XAI)!")

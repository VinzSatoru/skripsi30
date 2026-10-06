import docx
import re
import shutil

shutil.copy2("PROPOSAL BAB 1-3.docx", "temp_inspect_proposal.docx")
with open("temp_inspect_proposal.docx", "rb") as f:
    doc = docx.Document(f)

# Find start of DAFTAR PUSTAKA
start_dp = -1
for i, p in enumerate(doc.paragraphs):
    if "DAFTAR PUSTAKA" in p.text.upper():
        start_dp = i
        break

dp_entries = [p.text.strip() for p in doc.paragraphs[start_dp + 1:] if p.text.strip()]

print(f"Total DAFTAR PUSTAKA entries currently: {len(dp_entries)}")
for i, e in enumerate(dp_entries, 1):
    print(f"{i:2d}. {e[:85]}...")

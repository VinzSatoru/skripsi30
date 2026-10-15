import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# 8 Columns requested by user:
# 1. Peneliti (Tahun)
# 2. Judul Artikel
# 3. Dataset yang Digunakan
# 4. Algoritma dan Metode
# 5. Metrik
# 6. Hasil
# 7. Celah
# 8. Masalah yang Diselesaikan oleh Skripsi Ini

headers_8 = [
    "No",
    "Peneliti (Tahun)",
    "Judul Artikel",
    "Dataset yang Digunakan",
    "Algoritma dan Metode",
    "Metrik",
    "Hasil",
    "Celah (Research Gap)",
    "Masalah yang Diselesaikan oleh Skripsi Ini"
]

data_8 = [
    [
        "1",
        "Taher & Ayon (2024)",
        "Exploring Sleep Disorders: A Comparative Analysis of Machine Learning Algorithms on Sleep Health and Lifestyle Data",
        "Sleep Health and Lifestyle Dataset (374 instans, 13 fitur demografi dan gaya hidup)",
        "Random Forest, AdaBoost, Logistic Regression, Gradient Boosting Machine (GBM)",
        "Akurasi, Precision, Recall, F1-Score",
        "Gradient Boosting mencapai performa tertinggi: Akurasi 93,80%, Precision 94,10%, Recall 93,80%, dan F1-Score 93,90%.",
        "Model murni beroperasi sebagai black-box tanpa penjelasan kontribusi fitur individual, serta variabel kategorikal diubah via label encoding manual yang memicu bias ordinal buatan.",
        "Menerapkan algoritma CatBoost yang memproses fitur kategorikal secara native dan mengintegrasikan Tree-SHAP untuk membedah transparansi penyebab risiko secara personal per individu."
    ],
    [
        "2",
        "Lin et al. (2025)",
        "Evaluation of Sleep Quality and Influencing Factors Among Medical and Non-Medical Students Using Machine Learning Techniques in Fujian During the Public Health Emergencies",
        "Data survei kuesioner PSQI 20.645 mahasiswa di Fujian, Tiongkok",
        "Artificial Neural Network (ANN), Decision Tree, Gradient Boosting Trees, Naive Bayes",
        "Akurasi, Sensitivity, Specificity, AUC-ROC",
        "Model ensemble memetakan faktor risiko mahasiswa dengan Akurasi 86,40% dan AUC-ROC 0,891.",
        "Analisis hanya berlaku di tingkat agregat populasi (makro); tidak mampu memberikan penjelasan personal di tingkat individu (local explanation) untuk interaksi non-linier kebiasaan digital.",
        "Menganalisis interaksi non-linier fitur gaya hidup secara granular per individu menggunakan CatBoost dan visualisasi Tree-SHAP individual, bukan sekadar agregasi populasi."
    ],
    [
        "3",
        "Ha et al. (2023)",
        "Predicting the Risk of Sleep Disorders Using a Machine Learning–Based Simple Questionnaire: Development and Validation Study",
        "Data klinis 4.622 pasien dari 2 rumah sakit (Samsung Medical Center & Ewha Womans University Medical Center)",
        "Extreme Gradient Boosting (XGBoost) dengan seleksi fitur SHAP",
        "AUROC, Sensitivity, Specificity, Brier Score",
        "AUROC > 0,897 untuk ketiga gangguan tidur (OSA, COMISA, insomnia); menghasilkan kuesioner ringkas SLEEPS 9 item.",
        "Memerlukan encoding manual pada XGBoost; dataset berbasis kuesioner klinis dan data antropometrik statis (bukan kebiasaan harian modifiable); SHAP hanya untuk seleksi fitur awal, bukan visualisasi real-time per pasien.",
        "Menggunakan parameter gaya hidup harian yang dapat diubah (modifiable digital lifestyle), memproses data kategorikal secara native dengan CatBoost, dan menyajikan penjelasan SHAP interaktif real-time per pasien dalam sistem CDSS."
    ],
    [
        "4",
        "Xie et al. (2026)",
        "Identifying Influencing Factors Associated with Sleep Quality in Undergraduates Based on Partial Least Squares Regression and XGBoost",
        "Data kuesioner gaya hidup 1.842 mahasiswa sarjana",
        "Partial Least Squares Regression (PLSR) dan XGBoost",
        "R², RMSE, MAE, Variable Importance in Projection (VIP)",
        "XGBoost mengidentifikasi screen time, tingkat stres, dan konsumsi kafein sebagai prediktor utama kualitas tidur (R² = 0,782, RMSE = 0,412).",
        "One-hot encoding pada data kategori memecah fitur menjadi variabel dummy, memicu fragmentasi interpretasi; rekomendasi hanya bersifat umum tanpa takaran target perubahan kebiasaan yang terukur.",
        "Menghindari fragmentasi variabel dengan penanganan native CatBoost dan memetakan nilai SHAP langsung menjadi target intervensi perilaku kuantitatif (personalized actionable targets) seperti batas menit main HP dan cut-off kafein."
    ],
    [
        "5",
        "Das et al. (2025)",
        "Prevalence and Factors Associated with Insomnia Among Chronic Disease Patients in Bangladesh: A Machine Learning Study",
        "Data survei klinis 1.222 pasien penyakit kronis di Bangladesh",
        "KNN, Random Forest, SVM, GBM, XGBoost, dan CatBoost",
        "Akurasi, Precision, Recall, F1-Score, AUC-ROC",
        "CatBoost meraih performa tertinggi di antara seluruh model: Akurasi 71,67%, F1-Score 71,23%, dan AUC-ROC 77,27%.",
        "Target klasifikasi terbatas biner sederhana (Insomnia vs Non-Insomnia); sampel terbatas pada pasien penyakit kronis; evaluasi SHAP hanya beeswarm global; tidak diimplementasikan ke dalam prototipe CDSS interaktif.",
        "Mengklasifikasikan 4 tingkatan keparahan risiko secara bertingkat (Healthy, Mild, Moderate, Severe), pada populasi umum gaya hidup digital 100.000 instans, dilengkapi penjelasan lokal individual dan purwarupa sistem CDSS interaktif dengan simulasi What-If."
    ],
    [
        "6",
        "Srinivasu et al. (2024)",
        "XAI-Driven CatBoost Multi-Layer Perceptron Neural Network for Analyzing Breast Cancer",
        "Breast Cancer Wisconsin Diagnostic Dataset (WDBC) (569 instans, 30 fitur laboratorium)",
        "Hybrid CatBoost + Multi-Layer Perceptron (MLP) dengan seleksi fitur ANOVA dan Tree-SHAP",
        "Akurasi, Precision, Recall, F1-Score, Specificity",
        "Mencapai Akurasi 99,30%, Precision 99,20%, Recall 99,40%, F1-Score 99,30%.",
        "Diterapkan pada data laboratorium patologis onkologi dengan parameter fisik yang tidak dapat diubah (non-modifiable biological attributes); belum mengintegrasikan modul simulasi perubahan perilaku pasien.",
        "Mengadaptasi metodologi CatBoost dan SHAP ke ranah metrik perilaku gaya hidup digital yang modifiable, memungkinkan pengguna menguji perubahan kebiasaan secara simulatif pada antarmuka CDSS."
    ]
]

def format_cell(cell, text, bold=False, font_size=9, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def apply_table_styles(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header styling
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "F2F4F7")
    # Borders
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'bottom', 'insideH']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'CCCCCC')
        tblBorders.append(border)
    for border_name in ['left', 'right', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
    tblPr.append(tblBorders)

# 1. UPDATE PROPOSAL BAB 1-3 MASTER DOCX
def update_proposal_docx(file_path):
    doc = docx.Document(file_path)
    # Find table 2.1
    target_table = None
    for tbl in doc.tables:
        if len(tbl.rows) >= 6:
            target_table = tbl
            break
            
    if target_table:
        print(f"Found target table in {file_path}. Rows: {len(target_table.rows)}, Cols: {len(target_table.columns)}")
        # We need to replace target_table with an 9-column table (No + 8 columns)
        # Clear existing rows or build a new table right after it
        parent = target_table._tbl.getparent()
        idx = parent.index(target_table._tbl)
        
        # Create new table
        new_table = doc.add_table(rows=len(data_8) + 1, cols=len(headers_8))
        apply_table_styles(new_table)
        
        # Populate headers
        for c_idx, h in enumerate(headers_8):
            format_cell(new_table.rows[0].cells[c_idx], h, bold=True, font_size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            
        # Populate rows
        for r_idx, row_data in enumerate(data_8):
            for c_idx, val in enumerate(row_data):
                align = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                is_bold = (c_idx == 1)
                is_italic = (c_idx == 2)
                format_cell(new_table.rows[r_idx + 1].cells[c_idx], val, bold=is_bold, font_size=8.5, align=align, italic=is_italic)
                
        # Move new_table to idx position and remove old table
        parent.remove(target_table._tbl)
        parent.insert(idx, new_table._tbl)
        
        doc.save(file_path)
        print(f"Successfully updated Table 2.1 in {file_path} to 8-column format!")

# 2. UPDATE OUTLINE PROPOSAL MASTER DOCX
def update_outline_docx(file_path):
    doc = docx.Document(file_path)
    if len(doc.tables) >= 2:
        target_table = doc.tables[1] # Table 1 is SOTA
        parent = target_table._tbl.getparent()
        idx = parent.index(target_table._tbl)
        
        new_table = doc.add_table(rows=len(data_8) + 1, cols=len(headers_8))
        apply_table_styles(new_table)
        
        for c_idx, h in enumerate(headers_8):
            format_cell(new_table.rows[0].cells[c_idx], h, bold=True, font_size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            
        for r_idx, row_data in enumerate(data_8):
            for c_idx, val in enumerate(row_data):
                align = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                is_bold = (c_idx == 1)
                is_italic = (c_idx == 2)
                format_cell(new_table.rows[r_idx + 1].cells[c_idx], val, bold=is_bold, font_size=8.5, align=align, italic=is_italic)
                
        parent.remove(target_table._tbl)
        parent.insert(idx, new_table._tbl)
        
        doc.save(file_path)
        print(f"Successfully updated Table 1 in {file_path} to 8-column format!")

update_proposal_docx("01_Naskah_Proposal_Master/PROPOSAL_BAB_1-3_FINAL_SEMPURNA.docx")
update_outline_docx("01_Naskah_Proposal_Master/OUTLINE_PROPOSAL_REVISI_FINAL.docx")

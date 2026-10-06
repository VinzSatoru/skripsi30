import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('temp_work_proposal.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if any(h in txt for h in ["BAB I", "BAB II", "BAB III", "DAFTAR PUSTAKA", "1.1", "2.1", "3.1", "3.2", "3.3", "3.4", "3.5", "3.6"]):
        print(f"P[{i:3d}]: {txt[:70]}")

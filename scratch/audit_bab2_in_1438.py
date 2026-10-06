import docx
import re

doc = docx.Document("BAB 1_1438.docx")
banned = [
    r"\bsaya\b", r"\bkami\b", r"\bpenulis\b", r"\bkita\b",
    r"\bsangat canggih\b", r"\bluar biasa\b", r"\btebakan\b",
    r"\bbisa dibilang\b", r"\balat bernama\b", r"\bcanggih\b"
]

print("=== AUDIT TONE OF BAB 2 (P[31] to P[61]) ===")
found_issues = 0
for i in range(31, 62):
    p = doc.paragraphs[i]
    for b in banned:
        m = re.findall(b, p.text, re.IGNORECASE)
        if m:
            print(f"P[{i}] contains '{m}':\n  {p.text}\n")
            found_issues += 1

print(f"Total tone issues found in Bab 2: {found_issues}")

# Check word count of each paragraph in Bab 2
print("\n=== PARAGRAPH WORD COUNTS IN BAB 2 ===")
for i in range(31, 62):
    p = doc.paragraphs[i]
    words = len(p.text.split())
    if words > 10:
        print(f"P[{i}] ({p.style.name}): {words} words | {p.text[:60]}...")

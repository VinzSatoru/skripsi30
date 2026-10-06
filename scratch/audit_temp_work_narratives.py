import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('temp_work_proposal.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

narrative = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    words = len(txt.split())
    if words >= 50 and not txt.startswith("Tabel ") and not txt.startswith("Gambar ") and not txt.startswith("http") and not txt.startswith("Alqudah") and not txt.startswith("Amann"):
        narrative.append((i, words, txt[:50]))

print(f"Total narrative paragraphs: {len(narrative)}")
out_of_range = []
for idx, (p_num, w_cnt, snip) in enumerate(narrative, 1):
    status = "OK" if 120 <= w_cnt <= 130 else f"OUT OF RANGE ({w_cnt})"
    if status != "OK":
        out_of_range.append((p_num, w_cnt, snip))
    print(f"{idx:2d}. P[{p_num:3d}]: {w_cnt:3d} words [{status}] | {snip}...")

print(f"\nOut of range count: {len(out_of_range)}")

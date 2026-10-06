import docx
import shutil

target = "PROPOSAL BAB 1-3.docx"
shutil.copy2(target, target + ".bak")

doc = docx.Document(target)

cleaned_p = 0
for p in doc.paragraphs:
    if "\xa0" in p.text:
        cleaned_p += 1
        for r in p.runs:
            if "\xa0" in r.text:
                r.text = r.text.replace("\xa0", " ")

cleaned_t = 0
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if "\xa0" in cell.text:
                cleaned_t += 1
                for p in cell.paragraphs:
                    for r in p.runs:
                        if "\xa0" in r.text:
                            r.text = r.text.replace("\xa0", " ")

doc.save(target)
print(f"[OK] Cleaned \\xa0 in {cleaned_p} paragraphs and {cleaned_t} table cells.")

# Re-verify word counts
doc2 = docx.Document(target)
b1_words = [len(doc2.paragraphs[i].text.split()) for i in range(7, 13)]
print("Bab 1 Latar Belakang word counts:", b1_words)
assert all(120 <= w <= 130 for w in b1_words), "Bab 1 word count mismatch!"

b2_indices = [34, 35, 36, 38, 39, 41, 42, 43, 45, 46, 48, 49, 51, 52, 53, 58, 60]
b2_words = [len(doc2.paragraphs[i].text.split()) for i in b2_indices]
print("Bab 2 narrative word counts:", b2_words)
assert all(120 <= w <= 130 for w in b2_words), "Bab 2 word count mismatch!"

print("[ALL CHECKS PASSED PERFECTLY AFTER NBSP NORMALIZATION!]")

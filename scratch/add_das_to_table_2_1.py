import sys
import os
import shutil
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

base_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(base_dir)

src_file = os.path.join(parent_dir, "PROPOSAL_BAB_1-3_FINAL_SEMPURNA.docx")
target_master = os.path.join(parent_dir, "PROPOSAL BAB 1-3.docx")
target_revisi = os.path.join(parent_dir, "PROPOSAL_BAB_1-3_Revisi.docx")

with open(src_file, "rb") as f:
    doc = docx.Document(f)

def format_narrative(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.4)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

def format_caption(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(4)
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.bold = True

# 1. Update P[53] narrative paragraph
p53 = doc.paragraphs[53]
p53.text = (
    "Untuk menjembatani kesenjangan antara kebutuhan data kategorikal yang utuh dan transparansi inferensi klinis, "
    "kombinasi algoritma CatBoost dan SHAP mulai dieksplorasi dalam penelitian medis. Das et al. (2025) membuktikan "
    "keunggulan CatBoost dalam memprediksi insomnia, namun kajian tersebut masih terbatas pada klasifikasi biner pasien "
    "komorbid rumah sakit dengan visualisasi SHAP global. Di ranah onkologi, Srinivasu et al. (2024) mengintegrasikan "
    "CatBoost dan SHAP hingga meraih akurasi 99,3% dengan eksplanasi fitur yang jelas. Validasi keunggulan CatBoost atas "
    "data tabular medis juga diperkuat oleh Wolak et al. (2025). Meskipun demikian, penerapan sinergi CatBoost dan "
    "Tree-SHAP belum pernah dieksplorasi secara terpadu pada metrik gaya hidup digital untuk klasifikasi risiko multi-kelas. "
    "Mengisi celah riset inilah yang menjadi posisi kebaruan (novelty) utama penelitian ini, yakni menghadirkan sistem "
    "deteksi risiko yang akurat, transparan, dan dapat ditindaklanjuti."
)
format_narrative(p53)
print(f"[OK] P[53] updated: {len(p53.text.split())} words")

# 2. Update Table 2.1 caption
p54 = doc.paragraphs[54]
p54.text = "Tabel 2.1 Ringkasan Penelitian Terdahulu (State of the Art)"
format_caption(p54)
print(f"[OK] Table 2.1 caption updated to: {p54.text}")

# 3. Update Table 2.1 rows
t0 = doc.tables[0]
print(f"Current Table 2.1 rows: {len(t0.rows)}")

# Check if Das et al. is already in Table 2.1
das_found = any("Das et al." in r.cells[1].text for r in t0.rows)
if not das_found:
    # Add new row
    new_row = t0.add_row()
    # Copy Srinivasu to new row (Row 6)
    srinivasu_row = t0.rows[5]
    for c_idx in range(6):
        new_row.cells[c_idx].text = srinivasu_row.cells[c_idx].text
    new_row.cells[0].text = "6"
    
    # Overwrite Row 5 with Das et al. (2025)
    row5 = t0.rows[5]
    row5.cells[0].text = "5"
    row5.cells[1].text = "Das et al. (2025)"
    row5.cells[2].text = (
        "Prevalence and factors associated with insomnia among chronic disease patients in "
        "Bangladesh: A machine learning study (Nature and Science of Sleep)"
    )
    row5.cells[3].text = "CatBoost, XGBoost, Random Forest, SVM, LightGBM, SHAP"
    row5.cells[4].text = "CatBoost meraih akurasi tertinggi (71,67%) dan AUC 77,27% dalam memprediksi insomnia pada pasien komorbid."
    row5.cells[5].text = (
        "Klasifikasi terbatas pada target biner pasien rumah sakit; SHAP hanya disajikan pada agregat global tanpa "
        "visualisasi lokal per individu dan tanpa sistem CDSS terapan."
    )
    
    # Format all cells in Table 2.1
    for r_idx, r in enumerate(t0.rows):
        for c_idx, c in enumerate(r.cells):
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    if r_idx == 0:
                        run.bold = True
                    elif c_idx == 1:
                        run.bold = True
    print("[OK] Added Das et al. (2025) as Row 5 and shifted Srinivasu et al. (2024) to Row 6 in Table 2.1.")

# Save to work file
work_file = os.path.join(parent_dir, "temp_work_proposal.docx")
doc.save(work_file)
shutil.copy2(work_file, src_file)
shutil.copy2(work_file, target_master)
shutil.copy2(work_file, target_revisi)

print(f"[SUCCESS] All files successfully updated and synchronized!")

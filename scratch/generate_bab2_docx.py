import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

DOCX_OUT = r"c:\Users\ahmad\OneDrive\ドキュメント\skripsi\BAB_2_Draft_Final.docx"
MD_SRC = r"c:\Users\ahmad\OneDrive\ドキュメント\skripsi\BAB_2_Draft_Final.md"

def build_bab2_docx():
    doc = docx.Document()

    # Standard UNISNU Margins: Top 4cm, Left 4cm, Bottom 3cm, Right 3cm
    for section in doc.sections:
        section.top_margin = Cm(4)
        section.left_margin = Cm(4)
        section.bottom_margin = Cm(3)
        section.right_margin = Cm(3)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    # Style Helpers
    def set_font(run, name='Times New Roman', size=12, bold=False, italic=False, color=RGBColor(0,0,0)):
        run.font.name = name
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color

    def add_para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.5, first_indent=1.25):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if first_indent > 0:
            p.paragraph_format.first_line_indent = Cm(first_indent)
        r = p.add_run(text)
        set_font(r, size=12)
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(18)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        set_font(r, size=14, bold=True)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        set_font(r, size=12, bold=True)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        set_font(r, size=12, bold=True, italic=True)
        return p

    def format_table(table, col_widths_cm, alignments):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)
        
        for r_idx, row in enumerate(table.rows):
            is_header = (r_idx == 0)
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if is_header:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for c_idx, cell in enumerate(row.cells):
                if c_idx < len(col_widths_cm):
                    cell.width = Cm(col_widths_cm[c_idx])
                tcPr = cell._tc.get_or_add_tcPr()
                tcMar = OxmlElement('w:tcMar')
                for m, val in [('top', 120), ('bottom', 120), ('left', 140), ('right', 140)]:
                    node = OxmlElement(f'w:{m}')
                    node.set(qn('w:w'), str(val))
                    node.set(qn('w:type'), 'dxa')
                    tcMar.append(node)
                tcPr.append(tcMar)

                if is_header:
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F4F7"/>')
                    tcPr.append(shd)

                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    if c_idx < len(alignments):
                        p.alignment = alignments[c_idx]
                    for r in p.runs:
                        set_font(r, size=10, bold=is_header)

    # Read markdown and build document
    with open(MD_SRC, 'r', encoding='utf-8') as f:
        md_text = f.read()

    add_heading_1("BAB II\nLANDASAN TEORI DAN TINJAUAN PUSTAKA")

    # 2.1 Kajian Teori
    add_heading_2("2.1. Kajian Teori")
    add_para("Kajian teori memaparkan fondasi keilmuan, konsep fundamental, serta formulasi matematis yang melandasi perancangan sistem deteksi risiko gangguan tidur berbasis algoritma CatBoost dan metode Explainable Artificial Intelligence (XAI) Shapley Additive exPlanations (SHAP). Pembahasan teori mencakup karakteristik klinis gangguan tidur, pengaruh metrik gaya hidup digital, penanganan data tidak seimbang (class imbalance), arsitektur algoritma pembelajaran mesin berbasis gradient boosting, teori keadilan kontribusi fitur SHAP, hingga perancangan purwarupa sistem pendukung keputusan klinis (Clinical Decision Support System).")

    # 2.1.1 Karakteristik Klinis
    add_heading_3("2.1.1. Karakteristik Klinis Gangguan Tidur dan Kualitas Tidur")
    add_para("Tidur merupakan proses fisiologis esensial yang bersifat aktif dan berulang secara periodik untuk memulihkan fungsi homeostasis biologis, mengonsolidasikan memori kognitif, serta menjaga stabilitas sistem imun tubuh (Medic et al., 2017). Gangguan tidur (sleep disorders) didefinisikan sebagai deviasi atau kelainan pola tidur yang mengakibatkan penurunan kualitas, kuantitas, serta efisiensi tidur, yang berdampak langsung pada terganggunya fungsi fisiologis dan psikososial individu di siang hari. Berdasarkan International Classification of Sleep Disorders (ICSD-3), manifestasi gangguan tidur mencakup spektrum luas, mulai dari insomnia primer, gangguan pernapasan terkait tidur seperti Obstructive Sleep Apnea (OSA), gangguan ritme sirkadian, hingga hipersomnia (Ha et al., 2023).")
    add_para("Kualitas tidur diukur melalui keterpaduan parameter kuantitatif dan kualitatif, yang mencakup durasi total tidur (total sleep time), latensi awitan tidur (sleep onset latency), frekuensi terbangun di malam hari (wake after sleep onset), serta efisiensi tidur (sleep efficiency). Penurunan efisiensi tidur secara kronis berhubungan erat dengan peningkatan sitokin pro-inflamasi, resistensi insulin, peningkatan beban kardiovaskular, serta percepatan penurunan fungsi kognitif (Das et al., 2025). Selain durasi tidur statis, bukti klinis terkini menunjukkan bahwa variabilitas dan keteraturan jadwal tidur (sleep regularity) memegang peranan krusial sebagai prediktor risiko morbiditas dan mortalitas yang bahkan lebih kuat daripada durasi tidur itu sendiri (Windred et al., 2024).")
    add_para("Dalam konteks stratifikasi risiko klinis preventif, tingkat keparahan gangguan tidur dikelompokkan ke dalam empat tingkatan fungsional: Healthy (kondisi optimal tanpa indikasi patologis signifikan), Mild (gejala ringan yang dipicu oleh fluktuasi kebiasaan harian), Moderate (gangguan menengah yang mulai menurunkan produktivitas kognitif dan konsentrasi), serta Severe (gangguan berat dengan manifestasi klinis patologis yang memerlukan intervensi medis komprehensif). Klasifikasi berjenjang ini esensial untuk memfasilitasi tindakan triase dan skrining proaktif sebelum pasien berkembang menuju kondisi kronis.")

    # 2.1.2 Metrik Gaya Hidup Digital
    add_heading_3("2.1.2. Metrik Gaya Hidup Digital dan Pengaruhnya terhadap Ritme Sirkadian")
    add_para("Pergeseran pola interaksi masyarakat modern menuju era digital secara signifikan mengubah kebiasaan harian dan memunculkan fenomena stres psikologis baru yang berdampak negatif terhadap pola tidur. Jam biologis manusia diatur oleh ritme sirkadian yang berpusat pada suprachiasmatic nucleus (SCN) di hipotalamus, yang sangat peka terhadap paparan cahaya eksternal dan rutinitas harian. Paparan emisi cahaya biru (blue light) bergelombang pendek (450–480 nm) dari layar gawai (smartphone, laptop, tablet) menjelang waktu tidur menekan sekresi hormon melatonin oleh kelenjar pineal, sehingga menunda awitan tidur alami dan memperpendek fase tidur lelap (deep sleep) (Kumar et al., 2025).")
    add_para("Selain intervensi fotik fisiologis, aktivitas digital memicu fenomena ketergantungan psikologis yang dikenal sebagai nomophobia (no mobile phone phobia) dan technostress, yaitu kondisi tekanan mental akibat tuntutan konektivitas terus-menerus dan kebiasaan memeriksa gawai secara berulang di tempat tidur (Widayati, 2024; Jahrami, 2023). Kondisi ini meningkatkan aktivitas sistem saraf simpatik melalui sekresi hormon kortisol dan adrenalin, yang mengakibatkan kondisi hiperarousal (physiological hyperarousal) saat hendak beristirahat. Fenomena ini diperparah pada individu dengan tipe kronotipe malam (night owl) yang cenderung menunda jam tidur akibat aktivitas komputasi atau hiburan digital hingga larut malam.")
    add_para("Kombinasi antara paparan layar gawai yang intensif, beban kerja tinggi, asupan kafein berlebih, serta rendahnya tingkat aktivitas fisik harian (sedentary behavior) membentuk lingkaran umpan balik negatif terhadap kesehatan tidur (Henrich et al., 2021). Dampak kumulatif dari degradasi tidur akibat gaya hidup ini menimbulkan beban kerugian ekonomi yang nyata melalui penurunan produktivitas tenaga kerja, peningkatan absensi, dan risiko kesalahan fatal pada lingkungan profesional (Uzubuaku, 2023). Oleh karena itu, kuantifikasi metrik gaya hidup digital melalui variabel-variabel terukur menjadi instrumen penting untuk memodelkan risiko gangguan tidur secara presisi (Lin et al., 2025).")

    # 2.1.3 Pembelajaran Mesin Terawasi & Imbalance
    add_heading_3("2.1.3. Pembelajaran Mesin Terawasi dan Penanganan Ketidakseimbangan Data (Class Imbalance)")
    add_para("Pembelajaran mesin terawasi (supervised machine learning) merupakan paradigma komputasi di mana algoritma mempelajari fungsi pemetaan matematis dari ruang matriks fitur masukan X ke ruang label target luaran Y berdasarkan himpunan data latih berlabel. Pada domain medis, salah satu tantangan paling fundamental dalam pembelajaran terawasi adalah ketidakseimbangan sebaran kelas (class imbalance), di mana proporsi sampel populasi sehat (Healthy) mendominasi secara masif, sedangkan sampel pasien berisiko tinggi (Severe) hanya mencakup sebagian kecil dari total populasi (Taher & Ayon, 2024).")
    add_para("Algoritma pembelajaran mesin standar yang meminimalkan fungsi kerugian global tanpa penyesuaian akan mengalami fenomena Accuracy Paradox. Model akan cenderung memprediksi seluruh sampel ke dalam kelas mayoritas guna meraih angka akurasi komputasi yang tinggi, namun gagal mengenali pasien kelas minoritas yang justru memiliki risiko kritis secara klinis (false negative tinggi). Pendekatan konvensional untuk menangani masalah ini sering kali mengandalkan teknik manipulasi data (resampling), seperti random undersampling yang membuang informasi berharga dari kelas mayoritas, atau Synthetic Minority Over-sampling Technique (SMOTE) yang menghasilkan sampel sintetis di antara titik-titik data minoritas (Chicco & Jurman, 2020).")
    add_para("Akan tetapi, pada dataset tabular heterogen yang memadukan variabel numerik dan kategorikal, teknik SMOTE rentan menciptakan titik data sintetis yang tidak realistis secara klinis, merusak korelasi antar-variabel, serta meningkatkan risiko overfitting. Sebagai alternatif yang lebih kokoh dan mempertahankan struktur data asli, pendekatan Cost-Sensitive Learning diterapkan secara internal pada fungsi objektif optimasi model (El Chakik et al., 2026). Mekanisme ini menetapkan matriks biaya penalti kesalahan yang berbanding terbalik dengan frekuensi kemunculan kelas dalam data latih: wc = N / (K · Nc). Melalui penalti asimetris ini, setiap kekeliruan klasifikasi pada sampel kelas minoritas Severe diberikan bobot penalti hukuman belasan kali lipat lebih berat, sehingga memaksa algoritma memperluas batas keputusan untuk mengenali karakteristik pasien berisiko tinggi tanpa memanipulasi data riil (Chicco & Jurman, 2022).")

    # 2.1.4 Arsitektur Algoritma CatBoost
    add_heading_3("2.1.4. Arsitektur Algoritma CatBoost (Categorical Boosting)")
    add_para("CatBoost (Categorical Boosting) merupakan algoritma pembelajaran mesin mutakhir berbasis Gradient Boosted Decision Trees (GBDT) yang dikembangkan khusus untuk mengatasi keterbatasan algoritma boosting konvensional saat memproses data tabular berskala besar dan kaya fitur kategorikal (Prokhorenkova et al., 2018). Tidak seperti algoritma XGBoost atau LightGBM yang menggunakan pohon keputusan asimetris dengan pertumbuhan berbasis kedalaman (depth-wise) atau daun (leaf-wise), CatBoost memanfaatkan struktur pohon simetris (symmetric oblivious trees).")
    add_para("Pada oblivious decision tree, kriteria pemisahan (split criterion) yang sama diterapkan secara seragam pada seluruh simpul di tingkat kedalaman pohon yang identik. Struktur simetris ini berfungsi sebagai regularisasi alami yang mencegah pertumbuhan pohon yang terlampau dalam pada cabang tertentu, mempercepat evaluasi bitwise pada prosesor, serta menjamin stabilitas prediksi (Hancock & Khoshgoftaar, 2020; Li et al., 2025). Selain itu, inovasi Ordered Boosting pada CatBoost mengatasi bias kebocoran data (data leakage) dengan mensimulasikan proses pembelajaran sekuensial melalui permutasi acak observasi data (Wang et al., 2025; Srinivasu et al., 2024).")
    add_para("Keunggulan paling krusial dari CatBoost terletak pada mekanisme penanganan data kategorikal alami melalui Ordered Target Statistics (OTS). Pada data tabular medis, metode konvensional seperti One-Hot Encoding memecah variabel kategorikal menjadi puluhan variabel biner yang memicu lonjakan dimensi (curse of dimensionality) dan merusak penjelasan fitur (Chen et al., 2024). CatBoost mentransformasikan kategori menjadi nilai numerik kontinu secara dinamis berdasarkan nilai target historis ditambah pembobot prioritas global tanpa menimbulkan kebocoran informasi target masa depan (Alqudah et al., 2026).")

    # 2.1.5 Konsep XAI & SHAP
    add_heading_3("2.1.5. Konsep Explainable Artificial Intelligence (XAI) dan Teori Permainan Kooperatif SHAP")
    add_para("Penerapan sistem kecerdasan buatan dalam bidang medis menghadapi hambatan besar terkait sifat model kotak hitam (black-box model). Kendati algoritma pembelajaran mesin seperti CatBoost mampu mencapai akurasi klasifikasi yang tinggi, proses inferensi yang dihasilkan dari interaksi non-linier ratusan pohon keputusan tidak dapat dipahami secara intuitif oleh tenaga medis maupun pasien (Amann et al., 2020; Loh et al., 2024). Ketiadaan transparansi ini berisiko memicu penolakan adopsi teknologi oleh praktisi kesehatan.")
    add_para("Guna menjembatani kesenjangan tersebut, bidang Explainable Artificial Intelligence (XAI) menghadirkan metodologi untuk membuka proses penalaran internal model prediktif ke dalam representasi yang dapat diinterpretasikan oleh manusia. Pendekatan Shapley Additive exPlanations (SHAP) yang diperkenalkan oleh Lundberg & Lee (2017) berakar pada landasan teori permainan kooperatif (cooperative game theory) yang dirumuskan oleh Lloyd Shapley (1953). Prediksi model diposisikan sebagai hasil permainan (payout), sedangkan fitur masukan bertindak sebagai pemain (players) yang berkoalisi untuk menghasilkan keputusan tersebut.")
    add_para("Metode SHAP secara matematis menjamin empat aksioma keadilan fundamental: Efisiensi (jumlah nilai atribusi fitur ditambah baseline value persis sama dengan probabilitas keluaran model), Simetri (dua fitur dengan kontribusi marginal sama menerima atribusi identik), Pemain Nol (fitur non-kontributif menerima nilai nol), dan Monotonisitas. Modul TreeExplainer mengoptimalkan penghitungan nilai Shapley pada ansambel pohon menjadi kompleksitas waktu polinomial O(TLD^2), sehingga memungkinkan ekstraksi penjelasan global (summary beeswarm plot) dan lokal (waterfall plot) secara instan (Lundberg et al., 2020; Zhang et al., 2024; Huang et al., 2024).")

    # 2.1.6 CDSS
    add_heading_3("2.1.6. Sistem Pendukung Keputusan Klinis (Clinical Decision Support System / CDSS)")
    add_para("Sistem Pendukung Keputusan Klinis (Clinical Decision Support System / CDSS) merupakan aplikasi perangkat lunak berbasis teknologi informasi kesehatan yang dirancang untuk membantu tenaga medis maupun pasien dalam pengambilan keputusan klinis preventif maupun terapeutik (Naga Srinivasu et al., 2022). Evolusi CDSS modern bertransisi dari sistem pakar berbasis aturan statis (rule-based) menuju sistem cerdas adaptif yang didorong oleh algoritma pembelajaran mesin dan kapabilitas interpretabilitas XAI (Wahyudi et al., 2026).")
    add_para("Integrasi model prediktif CatBoost dan modul penjelasan SHAP ke dalam antarmuka purwarupa CDSS berbasis web interaktif menghadirkan skrining risiko non-invasif yang cepat, transparansi diagnostik per pasien, serta transformasi fitur risiko dominan menjadi rekomendasi perbaikan pola hidup (sleep hygiene interventions) yang dipersonalisasi dan berbasis bukti ilmiah (Chen et al., 2024; Rahman et al., 2025; Kaya, 2025; Putra & Hidayat, 2024).")

    # 2.2 SOTA
    add_heading_2("2.2. Kajian Hasil Penelitian yang Relevan (State of the Art)")
    add_para("Kajian hasil penelitian yang relevan menyajikan sintesis kritis terhadap studi-studi terdahulu yang berfokus pada klasifikasi gangguan tidur, penerapan algoritma gradient boosting, serta implementasi Explainable AI. Telaah ini bertujuan memetakan posisi keilmuan, mengidentifikasi keterbatasan yang ada, serta menegaskan kontribusi kebaruan (novelty) penelitian yang diusulkan.")
    
    add_heading_3("2.2.1. Sintesis Naratif Penelitian Terdahulu")
    add_para("Penelitian terdahulu mengenai klasifikasi gangguan tidur berbasis data kesehatan dan gaya hidup dilakukan oleh Taher & Ayon (2024) yang meraih akurasi 93,80% menggunakan Gradient Boosting. Namun, model bekerja murni sebagai kotak hitam tanpa transparansi faktor risiko individual. Lin et al. (2025) mengevaluasi 20.645 mahasiswa dan menemukan korelasi kuat antara gawai, kopi, dan penurunan tidur, tetapi analisisnya terbatas pada agregat populasi umum. Ha et al. (2023) menerapkan XGBoost dan SHAP pada 4.622 data kuesioner medis klinis dengan AUROC >0,897, namun memerlukan encoding manual dan SHAP hanya untuk seleksi fitur awal.")
    add_para("Kelemahan encoding ditegaskan oleh Xie et al. (2026) di mana penerapan One-Hot Encoding pada XGBoost memecah fitur kategorikal nominal sehingga penjelasan visual SHAP terfragmentasi dan membingungkan secara klinis. Terakhir, Das et al. (2025) pada Nature and Science of Sleep membuktikan CatBoost meraih performa tertinggi (akurasi 71,67%, F1 71,23%) dibanding 5 model lainnya dalam prediksi insomnia, namun terbatas pada klasifikasi biner pasien rumah sakit dengan SMOTE sintetis dan tanpa purwarupa sistem web operasional. Belum ada penelitian yang memadukan CatBoost native handling dan SHAP lokal untuk klasifikasi multi-kelas risiko gaya hidup secara terintegrasi.")

    # Tabel 2.1
    p_t21 = doc.add_paragraph()
    p_t21.paragraph_format.space_before = Pt(8)
    p_t21.paragraph_format.space_after = Pt(4)
    r_t21 = p_t21.add_run("Tabel 2.1 Matriks Komparasi Penelitian Terdahulu, Celah Riset (Research Gap), dan Solusi Kebaruan")
    set_font(r_t21, size=11, bold=True)

    tab_data = [
        ["No", "Peneliti & Tahun", "Metode / Algoritma", "Dataset & Variabel", "Temuan Utama", "Celah Riset (Research Gap)", "Solusi Penelitian Ini (Ahmad, 2026)"],
        ["1", "Taher & Ayon (2024)", "Random Forest, AdaBoost, Gradient Boosting", "Data gaya hidup dan tidur (Sleep Health & Lifestyle)", "Gradient Boosting meraih akurasi tertinggi 93,80%.", "Model black-box murni tanpa XAI; tidak menjelaskan faktor spesifik per individu.", "Menerapkan XAI-SHAP untuk membuka transparansi faktor risiko personal."],
        ["2", "Lin et al. (2025)", "ANN, Decision Tree, Naive Bayes", "20.645 data survei kuesioner mahasiswa", "Korelasi kuat antara durasi layar, kafein, dan kualitas tidur.", "Analisis agregat populasi umum; tidak ada personalisasi per pasien unik.", "Implementasi analisis SHAP lokal waterfall plot interaktif per individu."],
        ["3", "Ha et al. (2023)", "XGBoost + SHAP", "4.622 data kuesioner medis rumah sakit", "AUROC >0,897 untuk klasifikasi OSA, insomnia, dan COMISA.", "Wajib encoding manual, data kuesioner statis, SHAP hanya untuk seleksi fitur.", "CatBoost native categorical handling, fitur modifiable, dan SHAP real-time."],
        ["4", "Xie et al. (2026)", "PLS + XGBoost + SHAP", "Survei gaya hidup dan tidur mahasiswa", "Durasi tidur, stres, dan gawai prediktor utama.", "One-Hot Encoding memecah fitur nominal; visualisasi SHAP terfragmentasi.", "CatBoost ordered target statistics tanpa One-Hot, menjaga keutuhan semantik SHAP."],
        ["5", "Das et al. (2025)", "CatBoost, XGBoost, RF, SVM, GBM, KNN + SHAP", "1.222 data pasien kronis RS (Nature & Sci Sleep)", "CatBoost terbaik (Akurasi 71,67%, AUC 77,27%) pada insomnia.", "Target hanya biner, populasi pasien RS, eksplanasi global, tanpa web CDSS.", "Klasifikasi 4 level (Healthy s/d Severe) pada 100k data, SHAP dinamis, & Web CDSS."]
    ]
    t21 = doc.add_table(rows=len(tab_data), cols=7)
    for r_idx, row in enumerate(tab_data):
        for c_idx, val in enumerate(row):
            t21.rows[r_idx].cells[c_idx].paragraphs[0].text = val
    format_table(t21, [0.6, 2.0, 2.2, 2.2, 2.2, 2.4, 2.4],
                 [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    # 2.3 Kerangka Berpikir
    add_heading_2("2.3. Kerangka Berpikir (Conceptual Framework)")
    add_para("Kerangka berpikir penelitian ini dibangun berdasarkan alur logis pemecahan masalah kesehatan tidur di era digital melalui pendekatan pembelajaran mesin yang akuntabel. Permasalahan utama diawali dari tingginya prevalensi gangguan tidur yang dipicu oleh pola kebiasaan harian modern, seperti tingginya paparan layar gawai menjelang tidur, beban stres profesional, serta pola istirahat yang tidak teratur. Upaya deteksi dini yang ada saat ini menghadapi dua kendala utama: ketergantungan pada prosedur klinis invasif yang mahal serta kecenderungan model kecerdasan buatan konvensional yang bekerja sebagai kotak hitam (black-box) tanpa transparansi alasan medis.")
    add_para("Guna mengatasi permasalahan tersebut, kerangka penelitian ini mengintegrasikan data metrik gaya hidup digital berskala besar (100.000 rekaman data heterogen) ke dalam pipeline pemrosesan data mining CRISP-DM. Masalah ketidakseimbangan sebaran kelas target diatasi secara internal menggunakan pendekatan Cost-Sensitive Learning pembobotan penalti kelas terbalik, yang melatih algoritma CatBoost Classifier untuk mengenali pola pasien kelas minoritas berisiko tinggi (Severe) secara akurat tanpa merusak data asli melalui manipulasi oversampling. Algoritma CatBoost secara alami memproses atribut kategorikal melalui Ordered Target Statistics, sehingga menjaga keutuhan struktur matriks fitur.")
    add_para("Pada tahap interpretabilitas, modul Tree-SHAP diterapkan untuk membongkar mekanisme pengambilan keputusan model. Prinsip teori permainan kooperatif menjamin bahwa setiap variabel masukan menerima skor atribusi kontribusi marjinal yang adil, konsisten, dan memenuhi aksioma efisiensi aditif lokal. Keluaran inferensi model CatBoost beserta visualisasi waterfall plot SHAP kemudian diintegrasikan ke dalam antarmuka purwarupa Clinical Decision Support System (CDSS) berbasis web interaktif. Sistem ini tidak hanya menyajikan tingkat risiko tidur seseorang, melainkan juga memvisualisasikan faktor pemicu dominan dan menghasilkan rekomendasi perbaikan pola hidup (sleep hygiene) yang terpersonalisasi.")

    # 2.4 Pertanyaan Penelitian
    add_heading_2("2.4. Pertanyaan Penelitian (Research Questions)")
    add_para("Berdasarkan rumusan masalah dan kerangka berpikir yang telah dibangun, pertanyaan penelitian operasional yang akan dijawab dan dibuktikan secara empiris dalam penelitian ini adalah sebagai berikut:")
    
    p_q1 = doc.add_paragraph()
    p_q1.paragraph_format.left_indent = Cm(1.0)
    p_q1.paragraph_format.first_line_indent = Cm(-0.5)
    p_q1.paragraph_format.line_spacing = 1.5
    p_q1.paragraph_format.space_after = Pt(4)
    r1 = p_q1.add_run("1. ")
    set_font(r1, bold=True)
    r1_txt = p_q1.add_run("Apakah penerapan algoritma CatBoost dengan mekanisme Ordered Target Statistics dan penanganan ketidakseimbangan kelas berbasis Cost-Sensitive Learning mampu menghasilkan performa klasifikasi multikelas yang unggul, khususnya dalam mencapai nilai Recall kelas minoritas Severe di atas 85% dan Macro F1-Score di atas 85%?")
    set_font(r1_txt)

    p_q2 = doc.add_paragraph()
    p_q2.paragraph_format.left_indent = Cm(1.0)
    p_q2.paragraph_format.first_line_indent = Cm(-0.5)
    p_q2.paragraph_format.line_spacing = 1.5
    p_q2.paragraph_format.space_after = Pt(4)
    r2 = p_q2.add_run("2. ")
    set_font(r2, bold=True)
    r2_txt = p_q2.add_run("Bagaimana efektivitas metode XAI berbasis Tree-SHAP dalam mengidentifikasi dan memetakan variabel gaya hidup digital yang paling berkontribusi secara global pada tingkat populasi serta secara lokal pada tingkat individu tanpa kehilangan keutuhan semantik fitur?")
    set_font(r2_txt)

    p_q3 = doc.add_paragraph()
    p_q3.paragraph_format.left_indent = Cm(1.0)
    p_q3.paragraph_format.first_line_indent = Cm(-0.5)
    p_q3.paragraph_format.line_spacing = 1.5
    p_q3.paragraph_format.space_after = Pt(6)
    r3 = p_q3.add_run("3. ")
    set_font(r3, bold=True)
    r3_txt = p_q3.add_run("Bagaimana purwarupa Clinical Decision Support System (CDSS) berbasis web yang dirancang mampu menyajikan visualisasi probabilitas prediksi risiko dan penjelasan faktor pemicu dominan secara intuitif guna mendukung pengambilan keputusan intervensi klinis?")
    set_font(r3_txt)

    # DAFTAR PUSTAKA BAB II
    add_heading_2("DAFTAR PUSTAKA BAB II")
    
    # Extract references from markdown
    refs_text = md_text.split("## DAFTAR PUSTAKA BAB II")[1].strip()
    ref_entries = [r.strip() for r in refs_text.split("\n\n") if r.strip() and not r.startswith("Berikut adalah")]

    for ref in ref_entries:
        # Split citation text and status note
        lines = ref.split("\n")
        citation_line = lines[0].strip()
        status_line = lines[1].strip() if len(lines) > 1 else ""
        
        # Remove leading number like "1. "
        import re
        citation_clean = re.sub(r'^\d+\.\s*', '', citation_line)
        
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ref.paragraph_format.left_indent = Cm(1.0)
        p_ref.paragraph_format.first_line_indent = Cm(-1.0)
        p_ref.paragraph_format.line_spacing = 1.5
        p_ref.paragraph_format.space_after = Pt(4)
        
        r_cit = p_ref.add_run(citation_clean + "\n")
        set_font(r_cit, size=11)
        
        if status_line:
            r_stat = p_ref.add_run(status_line)
            set_font(r_stat, size=10, italic=True, color=RGBColor(80, 80, 80))

    try:
        doc.save(DOCX_OUT)
        print(f"[OK] Successfully built: {DOCX_OUT}")
    except PermissionError:
        alt_out = r"c:\Users\ahmad\OneDrive\ドキュメント\skripsi\BAB_2_Draft_Final_Revisi.docx"
        doc.save(alt_out)
        print(f"[NOTE] '{DOCX_OUT}' is open. Successfully saved to: {alt_out}")

if __name__ == "__main__":
    build_bab2_docx()

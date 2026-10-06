import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('OUTLINE_PROPOSAL_REVISI_FINAL.docx')

print("--- TABLE 0: FORMULIR PROPOSAL ---")
for r in doc.tables[0].rows:
    print([c.text.strip().replace('\n', ' ') for c in r.cells])

print("\n--- TABLE 1: RESEARCH GAP ---")
for r in doc.tables[1].rows:
    print([c.text.strip().replace('\n', ' ') for c in r.cells])

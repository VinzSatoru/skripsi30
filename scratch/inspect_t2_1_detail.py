import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('PROPOSAL_BAB_1-3_FINAL_SEMPURNA.docx')
t = doc.tables[0]

print(f"Table 2.1 Rows: {len(t.rows)}, Cols: {len(t.columns)}")
for r_idx, r in enumerate(t.rows):
    row_text = [c.text.strip().replace('\n', ' ') for c in r.cells]
    print(f"Row {r_idx}: {row_text[1]} | {row_text[3]} | {row_text[4][:30]}...")

# Check font settings in cells
for r_idx in [0, 1]:
    r = t.rows[r_idx]
    for c_idx, c in enumerate(r.cells):
        for p in c.paragraphs:
            for run in p.runs:
                print(f"R{r_idx}C{c_idx}: font={run.font.name}, size={run.font.size}, bold={run.bold}")

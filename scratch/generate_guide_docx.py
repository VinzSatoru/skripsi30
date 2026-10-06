import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_guide_docx():
    doc = docx.Document()

    # 1. Page Setup (A4, Standard 2.54 cm / 1 inch margins)
    for section in doc.sections:
        section.page_width = Inches(8.27)  # A4 Width
        section.page_height = Inches(11.69) # A4 Height
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 2. Helpers for Styling
    def set_font(run, name='Times New Roman', size=11, bold=False, italic=False, color=RGBColor(0,0,0)):
        run.font.name = name
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color

    def set_cell_shading(cell, color_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, border_color="CCCCCC", inside_color="E5E5E5"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{inside_color}"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{inside_color}"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_callout(text_title, text_body, border_hex="1F4E79", bg_hex="F2F6FA"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.columns[0].width = Inches(6.27)
        cell = tbl.rows[0].cells[0]
        set_cell_shading(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        # Left thick border
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        cell._tc.get_or_add_tcPr().append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run(text_title + "\n"), bold=True, size=10.5, color=RGBColor(31, 78, 121))
        set_font(p.add_run(text_body), size=10, italic=False)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # COVER / HEADER
    # -------------------------------------------------------------
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p_header.add_run("PANDUAN STRATEGIS PENGAJUAN JUDUL SKRIPSI,\nANTISIPASI TANYA-JAWAB KRITIS, DAN GLOSARIUM TEKNIS\n")
    set_font(r0, size=13.5, bold=True, color=RGBColor(31, 78, 121))
    r1 = p_header.add_run("Program Studi S1 Teknik Informatika — Fakultas Sains dan Teknologi\nUniversitas Islam Nahdlatul Ulama (UNISNU) Jepara\n")
    set_font(r1, size=11, bold=True, color=RGBColor(80, 80, 80))

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.line_spacing = 1.15
    p_meta.paragraph_format.space_after = Pt(12)
    set_font(p_meta.add_run("Nama Mahasiswa : "), bold=True, size=10)
    set_font(p_meta.add_run("Ahmad Novian Dzulfanni\n"), size=10)
    set_font(p_meta.add_run("NIM            : "), bold=True, size=10)
    set_font(p_meta.add_run("231240001438\n"), size=10)
    set_font(p_meta.add_run("Judul Skripsi  : "), bold=True, size=10)
    set_font(p_meta.add_run("Penerapan Metode XAI-SHAP pada Algoritma CatBoost untuk Klasifikasi Faktor Risiko Gangguan Tidur Berbasis Metrik Gaya Hidup Digital\n"), bold=True, size=10, color=RGBColor(0, 102, 204))

    # -------------------------------------------------------------
    # BAGIAN 1: STRATEGI & DRAFT KALIMAT PENGAJUAN JUDUL
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    set_font(h1.add_run("BAGIAN 1: STRATEGI & DRAFT KALIMAT PENGAJUAN JUDUL"), size=12.5, bold=True, color=RGBColor(31, 78, 121))

    p1 = doc.add_paragraph()
    p1.paragraph_format.line_spacing = 1.15
    set_font(p1.add_run("Saat seorang mahasiswa mengajukan judul skripsi ke Ketua Program Studi (Kaprodi) atau calon Dosen Pembimbing, dosen pada dasarnya menerapkan 4 filter evaluasi bawah sadar:\n"), size=10.5)
    set_font(p1.add_run("1. Apakah masalahnya nyata dan relevan di ranah informatika modern?\n2. Apakah mahasiswa paham betul alasan pemilihan algoritmanya (bukan sekadar ikut-ikutan)?\n3. Apakah datanya sudah siap, valid, dan cukup skalanya (menghindari mahasiswa macet di tengah jalan)?\n4. Apakah ada produk luaran yang konkret (tidak hanya skrip di komputer)?\n"), size=10, italic=True)

    add_callout(
        "OPSI A: NASKAH PENGUCAPAN LANGSUNG / TATAP MUKA (ELEVATOR PITCH ±90 DETIK)",
        "Gunakan intonasi yang tenang, tatap mata dosen, dan bicara dengan tempo teratur:\n\n"
        "\"Selamat pagi/siang, Bapak/Ibu [Nama Kaprodi/Dosen]. Mohon izin mengonsultasikan rencana judul skripsi saya.\n\n"
        "Judul yang saya ajukan adalah: 'Penerapan Metode XAI-SHAP pada Algoritma CatBoost untuk Klasifikasi Faktor Risiko Gangguan Tidur Berbasis Metrik Gaya Hidup Digital'.\n\n"
        "Latar belakang saya mengangkat topik ini adalah tingginya tren gangguan tidur di era modern yang sangat dipengaruhi oleh kebiasaan harian (seperti durasi layar gawai, beban kerja, dan stres). Berdasarkan kajian terhadap 5 jurnal internasional bereputasi terkini—termasuk jurnal Nature and Science of Sleep tahun 2025—penelitian deteksi risiko tidur saat ini memiliki dua keterbatasan utama:\n"
        "1. Banyak algoritma konvensional (seperti XGBoost atau Random Forest) yang mengolah data kategorikal gaya hidup menggunakan One-Hot Encoding sehingga memecah keutuhan fitur dan menimbulkan ledakan dimensi data.\n"
        "2. Model yang dibangun umumnya beroperasi sebagai kotak hitam (black-box), sehingga tidak dapat menjelaskan faktor risiko personal pada masing-masing individu secara transparan.\n\n"
        "Oleh karena itu, pada penelitian ini saya mengusulkan algoritma CatBoost karena memiliki keunggulan native handling categorical data tanpa distorsi One-Hot Encoding, yang kemudian diintegrasikan dengan metode Explainable AI berbasis SHAP untuk menghasilkan penjelasan kontribusi faktor risiko secara lokal per individu (waterfall plot) yang adil dan matematis.\n\n"
        "Untuk kesiapan teknis, dataset sudah sangat siap sebanyak 100.000 rekaman data gaya hidup dengan klasifikasi 4 level risiko (Healthy, Mild, Moderate, Severe). Luaran akhirnya tidak hanya model komputasi, melainkan diimplementasikan langsung ke dalam purwarupa Web Clinical Decision Support System (CDSS) yang interaktif.\n\n"
        "Berkas outline formulir resmi dan pemetaan 5 Research GAP-nya sudah saya susun rapi, Pak/Bu. Mohon arahan dan persetujuannya agar saya dapat melanjutkan ke tahap sidang proposal. Terima kasih, Pak/Bu.\""
    )

    add_callout(
        "OPSI B: NASKAH PESAN DARING (FORMAT WHATSAPP / EMAIL FORMAL)",
        "Assalamu’alaikum Warahmatullahi Wabarakatuh / Selamat Pagi Bapak/Ibu [Nama Kaprodi],\n\n"
        "Mohon maaf mengganggu waktu Bapak/Ibu. Saya Ahmad Novian Dzulfanni (NIM: 231240001438), mahasiswa S1 Teknik Informatika peminatan Sains Data dan Machine Learning.\n\n"
        "Izin menyampaikan pengajuan rencana judul dan topik skripsi untuk dimohonkan telaah serta persetujuannya:\n\n"
        "📌 Usulan Judul:\n"
        "\"Penerapan Metode XAI-SHAP pada Algoritma CatBoost untuk Klasifikasi Faktor Risiko Gangguan Tidur Berbasis Metrik Gaya Hidup Digital\"\n\n"
        "📌 Urgensi Masalah:\n"
        "Peningkatan gangguan tidur era digital sangat dipicu oleh faktor gaya hidup yang dapat dimodifikasi (screen time, kafein, stres). Namun, model machine learning medis saat ini umumnya bersifat kotak hitam (black-box) sehingga tidak mampu menjelaskan alasan di balik keputusan prediksi per individu.\n\n"
        "📌 Solusi Metodologi & Kebaruan (Novelty):\n"
        "1. Menggunakan CatBoost yang memproses data gaya hidup heterogen secara native tanpa merusak data lewat One-Hot Encoding.\n"
        "2. Menerapkan XAI-SHAP untuk membedah kontribusi faktor risiko secara transparan dan adil bagi setiap individu (waterfall plot real-time).\n"
        "3. Menutup celah riset dari 5 jurnal internasional terkini (termasuk Das et al., Nature and Science of Sleep 2025) dengan menghadirkan klasifikasi multi-kelas 4 tingkat risiko.\n\n"
        "📌 Kesiapan Riset & Luaran:\n"
        "- Dataset: 100.000 data gaya hidup digital terstruktur.\n"
        "- Luaran Akhir: Model klasifikasi CatBoost-SHAP + Purwarupa Web Clinical Decision Support System (CDSS) interaktif.\n\n"
        "Bersama pesan ini, saya lampirkan dokumen formulir Outline Proposal resmi (OUTLINE_PROPOSAL_REVISI_FINAL.docx) yang telah memuat tabel matriks Research GAP lengkap sebagai bahan pertimbangan Bapak/Ibu.\n\n"
        "Besar harapan saya topik ini dapat disetujui untuk melangkah ke tahap seminar proposal. Terima kasih banyak atas waktu dan bimbingan Bapak/Ibu.\n\n"
        "Wassalamu’alaikum Warahmatullahi Wabarakatuh.\n"
        "Ahmad Novian Dzulfanni (231240001438)"
    )

    # -------------------------------------------------------------
    # BAGIAN 2: BANK 10 PERTANYAAN JEBAKAN & JAWABAN SKAKMAT
    # -------------------------------------------------------------
    doc.add_page_break()
    h2 = doc.add_heading(level=1)
    set_font(h2.add_run("BAGIAN 2: BANK PERTANYAAN JEBAKAN DOSEN & JAWABAN SKAKMAT ILMIAH"), size=12.5, bold=True, color=RGBColor(31, 78, 121))

    qa_list = [
        ("Pertanyaan 1: Kenapa harus pakai CatBoost? Kenapa tidak Random Forest, SVM, atau XGBoost yang lebih populer?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Data gaya hidup memadukan fitur numerik dan kategorikal (seperti jenis pekerjaan, gender, dan kondisi mental). Algoritma konvensional seperti XGBoost dan Random Forest mewajibkan One-Hot Encoding yang memecah satu kolom kategori menjadi puluhan kolom biner artifisial, sehingga memicu ledakan dimensi data dan membuat penjelasan SHAP terpecah-pecah.\n"
         "CatBoost memiliki mekanisme bawaan Ordered Target Statistics yang memproses data kategorikal secara native tanpa One-Hot Encoding, serta menggunakan Symmetric Oblivious Trees yang bertindak sebagai regularisasi alami pencegah overfitting dan mempercepat komputasi inferensi.\""),

        ("Pertanyaan 2: Kenapa butuh Explainable AI (SHAP)? Bukankah di Machine Learning yang penting akurasinya tinggi?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Di bidang kesehatan, akurasi tinggi saja belum cukup, Pak/Bu. Model ensemble seperti CatBoost terdiri atas ratusan pohon keputusan yang rumit sehingga bersifat black-box. Jika sistem mendeteksi seorang pasien berisiko tinggi tanpa penjelasan, dokter dan pasien tidak akan percaya dan tidak tahu apa yang harus diperbaiki.\n"
         "SHAP yang berbasis cooperative game theory (Shapley Values) menjamin perhitungan kontribusi setiap variabel secara adil, konsisten, dan aditif. Melalui analisis lokal (waterfall plot), pasien bisa tahu pasti bahwa risikonya tinggi akibat screen time 8 jam dan stres level 7, sehingga rekomendasi sleep hygiene yang diberikan tepat sasaran.\""),

        ("Pertanyaan 3: Datanya dari mana? Jangan-jangan data sekunder sedikit atau data main-main?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Data penelitian kami berskala besar dan terstruktur, yaitu 100.000 rekaman data gaya hidup digital yang mencakup 28 fitur relevan (durasi tidur, detak jantung, screen time, kafein, langkah harian, hingga kondisi mental). Data ini bebas dari nilai hilang (missing values) dan memiliki sebaran target 4 kelas (Healthy, Mild, Moderate, Severe) yang sangat representatif untuk simulasi sistem deteksi dini.\""),

        ("Pertanyaan 4: Fokus skripsimu ini di mana? Menguji algoritma Machine Learning atau membuat aplikasi Web?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Fokus utama penelitian saya 80% berada pada eksperimen Sains Data dan Machine Learning terapan dengan mengadopsi standar industri CRISP-DM secara utuh.\n"
         "Tahap 1 sampai 5 (Business Understanding hingga Evaluation) berfokus penuh pada perancangan CatBoost, penanganan imbalanced data, dan validasi matematis XAI-SHAP.\n"
         "Sedangkan purwarupa Web CDSS adalah tahap ke-6 (Deployment) sebagai Proof of Concept (pembuktian terapan) agar kecerdasan model tidak berhenti di notebook, melainkan bisa dioperasikan secara nyata oleh dokter dan pasien. Jadi sistem web adalah media hilirisasi dari otak model Machine Learning-nya.\""),

        ("Pertanyaan 5: Data kamu sangat tidak seimbang (Severe cuma 4%), bagaimana cara mengatasinya dan kenapa tidak pakai SMOTE atau Undersampling?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Kami menerapkan pendekatan Cost-Sensitive Learning bawaan CatBoost melalui auto_class_weights='Balanced'. Kami tidak menggunakan Undersampling karena tidak ingin membuang 50.000 data riil orang sehat, dan tidak menggunakan SMOTE karena pembuatan data sintetis pada kombinasi data kategorikal berisiko merusak relasi logis variabel medis dan memicu overfitting.\n"
         "Secara matematis, mekanisme Balanced Loss memberikan penalti hukuman 13 kali lipat lebih berat pada kelas Severe (bobot 6,15) dibanding kelas Healthy (bobot 0,46), sehingga algoritma dipaksa secara matematis untuk memprioritaskan pendeteksian pasien risiko parah tanpa memanipulasi data aslinya.\""),

        ("Pertanyaan 6: Kamu training data berapa kali sehingga dapat akurasi yang sesuai?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Proses training dilakukan melalui dua aspek terstruktur:\n"
         "1. Secara Eksperimen Metodologis: Kami menguji skenario hyperparameter tuning secara sistematis membandingkan learning rate (0.01 s/d 0.05), depth (4, 6, 8), dan skema pembobotan kelas untuk memperoleh kombinasi parameter terbaik.\n"
         "2. Secara Algoritma Internal: CatBoost dilatih dengan batas maksimum 1.000 iterasi pohon, tetapi dikontrol oleh Early Stopping sebanyak 50 putaran evaluasi terhadap 20.000 data uji validasi. Begitu performa validasi tidak lagi membaik dalam 50 putaran, proses pelatihan langsung berhenti otomatis di titik konvergensi optimal untuk mencegah overfitting.\""),

        ("Pertanyaan 7: Pada dataset ini, apa atribut yang menjadi X, Y, dan Z?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Dalam konvensi supervised learning, hanya terdapat variabel X dan Y:\n"
         "- Variabel Y (Target): 1 kolom yaitu sleep_disorder_risk yang terdiri dari 4 kelas klasifikasi (Healthy, Mild, Moderate, Severe).\n"
         "- Variabel X (Fitur Prediktor): 27–28 variabel masukan yang mencakup parameter gaya hidup digital, metrik tidur, kondisi fisiologis, dan lingkungan.\n"
         "- Variabel Z: Tidak ada variabel Z dalam dataset asli. Namun jika dikaitkan dengan metode XAI, nilai Z merepresentasikan Matriks Kontribusi SHAP yang dihitung untuk menjelaskan pengaruh masing-masing fitur X terhadap keputusan Y.\n"
         "Selain itu, atribut person_id, country, dan day_type sengaja dibuang (drop) untuk mencegah bias geografis dan menghapus atribut tanpa nilai prediktif.\""),

        ("Pertanyaan 8: Apa kebaruan (novelty) risetmu dibanding penelitian yang sudah ada sebelumnya?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Penelitian ini mengisi celah dari 5 jurnal internasional bereputasi terkini yang 100% homogen di domain gangguan tidur:\n"
         "1. Taher & Ayon (2024) hanya model black-box tanpa penjelasan faktor risiko personal.\n"
         "2. Lin et al. (2025) hanya menganalisis faktor tidur pada tingkat populasi umum secara agregat.\n"
         "3. Ha et al. (2023) menggunakan data kuesioner statis dan SHAP hanya untuk seleksi fitur awal.\n"
         "4. Chen et al. (2024) memakai XGBoost dengan One-Hot Encoding yang membuat eksplanasi SHAP terfragmentasi.\n"
         "5. Das et al. (2025 di Nature and Science of Sleep) membuktikan CatBoost-SHAP unggul, namun terbatas pada target biner pasien rumah sakit dan visualisasi global saja.\n"
         "Kebaruan penelitian ini adalah mengintegrasikan CatBoost native dengan SHAP waterfall lokal per individu untuk klasifikasi multi-kelas 4 level risiko pada 100.000 data gaya hidup serta diimplementasikan ke Web CDSS interaktif.\""),

        ("Pertanyaan 9: Kenapa kamu mengevaluasi performa menggunakan Macro F1-Score dan Recall Severe, bukan cuma Akurasi?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Karena dataset kami mengalami ketidakseimbangan kelas (Healthy 54% vs Severe 4%), evaluasi hanya dengan akurasi akan menimbulkan fenomena Accuracy Paradox di mana model terlihat pintar padahal hanya menebak kelas mayoritas.\n"
         "Oleh karena itu, kami memvalidasi performa dengan:\n"
         "1. Recall Severe sebesar 89%: Membuktikan bahwa dari seluruh pasien yang benar-benar sakit parah, hampir 90% berhasil dideteksi dengan tepat tanpa kecolongan (minim False Negative).\n"
         "2. Macro F1-Score sebesar 89%: Memberikan bobot hak suara yang sama rata (masing-masing 25%) pada tiap kelas tanpa memandang jumlah data, membuktikan keandalan model merata di semua level risiko.\""),

        ("Pertanyaan 10: Bagaimana kaitan purwarupa Web ini dengan dunia klinis nyata?",
         "Alasan Ilmiah Skakmat:\n"
         "\"Web CDSS ini berfungsi sebagai jembatan komunikasi antara pasien dan tenaga medis. Pasien dapat melakukan skrining mandiri secara cepat berdasarkan kebiasaan hariannya, sementara dokter mendapatkan visualisasi grafik waterfall yang membedah kebiasaan buruk spesifik pasien tersebut.\n"
         "Dengan demikian, rekomendasi intervensi medis yang diberikan bersifat personal, tepat sasaran, dan didukung bukti komputasi yang transparan, bukan sekadar tebakan umum.\"")
    ]

    for q_title, q_ans in qa_list:
        add_callout(q_title, q_ans, border_hex="2B4C7E", bg_hex="F9FBFD")

    # -------------------------------------------------------------
    # BAGIAN 3: GLOSARIUM LENGKAP ISTILAH ASING / TEKNIS
    # -------------------------------------------------------------
    doc.add_page_break()
    h3 = doc.add_heading(level=1)
    set_font(h3.add_run("BAGIAN 3: GLOSARIUM LENGKAP ISTILAH ASING & TEKNIS"), size=12.5, bold=True, color=RGBColor(31, 78, 121))

    p3_desc = doc.add_paragraph()
    p3_desc.paragraph_format.line_spacing = 1.15
    set_font(p3_desc.add_run("Tabel ini memuat istilah-istilah teknis penting yang sering muncul dalam skripsi ini, dilengkapi dengan definisi ilmiah formal dan analogi 'bahasa manusia' agar Anda dapat menjelaskannya dengan lugas dan mudah dipahami:"), size=10, italic=True)

    glossary_data = [
        ("Explainable AI (XAI)",
         "Kecerdasan Buatan Terjelaskan",
         "Cabang ilmu kecerdasan buatan yang berfokus pada pengembangan metode dan teknik agar keluaran keputusan algoritma pembelajaran mesin dapat dipahami, diinterpretasikan, dan dipercaya oleh manusia secara transparan.",
         "Membuka 'kap mesin' kecerdasan buatan agar kita tahu alasan logis mengapa sistem mengeluarkan diagnosis tertentu, bukan cuma menerima hasilnya mentah-mentah."),

        ("SHAP (Shapley Additive exPlanations)",
         "Metode SHAP (Lundberg & Lee, 2017)",
         "Kerangka kerja interpretasi model berbasis teori permainan kooperatif (Shapley values) yang menghitung kontribusi marjinal setiap fitur input terhadap hasil prediksi akhir model secara aditif, adil, dan konsisten.",
         "Seperti menilai kontribusi masing-masing pemain sepak bola dalam sebuah tim untuk menentukan siapa yang paling berjasa mencetak gol kemenangan."),

        ("CatBoost (Categorical Boosting)",
         "Algoritma Pembelajaran Mesin (Prokhorenkova et al., Yandex)",
         "Algoritma supervised learning berbasis gradient boosting decision trees yang dirancang khusus untuk menangani fitur kategorikal secara native menggunakan teknik Ordered Target Statistics.",
         "Algoritma pohon keputusan canggih yang sangat pintar mengolah data teks/kategori (seperti profesi atau gender) tanpa perlu diubah manual jadi angka biner."),

        ("Ordered Target Statistics",
         "Statistik Target Terurut",
         "Mekanisme pengkodean fitur kategorikal pada CatBoost yang menghitung nilai statistik target secara bertahap berdasarkan urutan acak data untuk mencegah target leakage dan overfitting.",
         "Cara cerdas mengubah kategori menjadi nilai bobot angka tanpa 'mengintip' jawaban data pengujian masa depan."),

        ("Symmetric Oblivious Trees",
         "Pohon Keputusan Simetris",
         "Arsitektur pohon keputusan biner yang menerapkan kriteria pembagian (split) yang persis sama di seluruh simpul pada level kedalaman yang sama.",
         "Pohon keputusan yang bentuk cabangnya seimbang dan rapi kiri-kanan, sehingga proses eksekusi prediksi sangat cepat dan tahan terhadap hafalan data berlebih."),

        ("Black-Box Model",
         "Model Kotak Hitam",
         "Model pembelajaran mesin yang struktur internal perhitungannya sangat rumit (seperti ratusan ensemble pohon atau jaringan saraf tiruan dalam) sehingga proses penalaran keputusannya tidak bisa dilihat secara kasat mata oleh manusia.",
         "Sistem yang menerima input lalu mengeluarkan output, tapi proses di dalamnya seperti ruang gelap gulita yang tidak ada jendelanya."),

        ("Imbalanced Data",
         "Ketidakseimbangan Kelas",
         "Kondisi pada dataset klasifikasi di mana jumlah sampel pada satu atau lebih kelas target jauh lebih banyak dibandingkan kelas lainnya (misal Healthy 54% vs Severe 4%).",
         "Kondisi ketika populasi orang sehat jauh melimpah ruah dibanding orang yang sakit parah dalam data survei."),

        ("Accuracy Paradox",
         "Paradoks Akurasi",
         "Kondisi menipu di mana model klasifikasi meraih skor akurasi persentase yang sangat tinggi pada data tidak seimbang, padahal model hanya menebak kelas mayoritas dan gagal total mendeteksi kelas minoritas.",
         "Seorang dokter yang selalu menebak 'kamu sehat' ke semua pasien. Akurasinya 90% karena 90 orang memang sehat, tapi dia membiarkan 10 orang sakit parah meninggal tanpa diobati."),

        ("Cost-Sensitive Learning",
         "Pembelajaran Peka Biaya (auto_class_weights='Balanced')",
         "Pendekatan pelatihan algoritma yang memberikan bobot penalti kerugian (loss) yang berbeda terhadap kesalahan klasifikasi, di mana kesalahan pada kelas minoritas dihukum jauh lebih berat daripada kelas mayoritas.",
         "Aturan ujian di mana salah menjawab soal orang sakit parah dikurangi 6 poin, sedangkan salah soal orang sehat hanya dikurangi 0,4 poin. Algoritma otomatis fokus belajar soal bernilai tinggi."),

        ("Stratified Sampling 80:20",
         "Pembagian Data Berstrata",
         "Teknik membagi dataset menjadi data latih dan data uji dengan mempertahankan persentase proporsi masing-masing kelas target secara identik pada kedua himpunan bagian.",
         "Memotong kue lapis sedemikian rupa sehingga potongan latihan dan potongan pengujian sama-sama memiliki proporsi rasa yang seimbang persis dengan kue aslinya."),

        ("Macro-averaged F1-Score",
         "Rata-Rata F1 Makro",
         "Metrik evaluasi yang menghitung nilai F1-score untuk masing-masing kelas secara individual, kemudian merata-ratakannya dengan bobot hak suara yang sama rata (unweighted) tanpa dipengaruhi oleh banyaknya sampel tiap kelas.",
         "Sistem pemilu di mana suara pulau kecil yang penduduknya sedikit memiliki kekuatan suara yang setara dengan pulau besar yang padat penduduk."),

        ("Recall (Sensitivitas)",
         "Daya Ingat / Deteksi Sensitif",
         "Metrik yang mengukur proporsi kasus positif aktual yang berhasil diidentifikasi secara tepat oleh model: TP / (TP + FN).",
         "Dari 100 orang yang benar-benar sakit parah, berapa orang yang berhasil dijaring oleh sistem dan tidak lolos dari pemeriksaan."),

        ("Precision (Presisi)",
         "Ketepatan Tebakan",
         "Metrik yang mengukur seberapa banyak tebakan positif model yang memang terbukti benar di kenyataan: TP / (TP + FP).",
         "Dari 100 orang yang divonis sakit oleh sistem, berapa orang yang benar-benar sakit sesungguhnya."),

        ("Confusion Matrix",
         "Matriks Kebingungan 4x4",
         "Tabel kontingensi yang memetakan perbandingan antara kelas aktual dengan kelas prediksi model untuk melihat rincian jumlah tebakan yang tepat dan arah letak kesalahan klasifikasinya.",
         "Tabel rekapitulasi nilai ujian yang menunjukkan di bagian mana sistem menjawab benar dan di bagian mana sistem salah menebak."),

        ("Waterfall Plot",
         "Grafik Air Terjun Lokal SHAP",
         "Visualisasi penjelasan lokal yang membedah bagaimana nilai dasar (base value) prediksi bergeser naik atau turun dipengaruhi oleh kontribusi spesifik masing-masing variabel input milik seorang pasien.",
         "Struk rincian tagihan belanja yang menjelaskan dari saldo awal, item apa saja yang menambah biaya dan item apa yang memberi diskon hingga keluar total harga akhir."),

        ("Summary Plot / Beeswarm",
         "Grafik Sebaran Global SHAP",
         "Visualisasi agregat yang merangkum tingkat kepentingan (importance) seluruh fitur dan arah pengaruh nilai variabelnya terhadap prediksi pada tingkat keseluruhan populasi data.",
         "Peta survei nasional yang memperlihatkan faktor apa saja yang secara umum paling sering membuat masyarakat Indonesia sulit tidur lelap."),

        ("Data Leakage / Target Leakage",
         "Kebocoran Informasi Fitur",
         "Kondisi keliru di mana variabel yang dimasukkan ke dalam model memuat informasi yang seharusnya baru diketahui setelah target terjadi, sehingga menghasilkan akurasi palsu yang tidak realistis.",
         "Mahasiswa yang bisa menjawab soal ujian karena lembar kunci jawabannya tidak sengaja tertempel di balik kertas soal."),

        ("Early Stopping",
         "Penghentian Dini Pelatihan",
         "Mekanisme regularisasi yang memantau performa model pada set data validasi di setiap iterasi dan menghentikan proses pelatihan secara otomatis apabila skor validasi tidak lagi menunjukkan perbaikan dalam sejumlah putaran tertentu.",
         "Berhenti memasak sayur tepat saat kuah sudah mendidih matang sempurna, sehingga masakan tidak gosong atau terlalu lembek."),

        ("CRISP-DM",
         "Cross-Industry Standard Process for Data Mining",
         "Standar metodologi proses penambangan data industri yang terdiri dari 6 tahapan berulang: Business Understanding, Data Understanding, Data Preparation, Modeling, Evaluation, dan Deployment.",
         "Buku panduan resep kerja resmi para ilmuwan data dunia agar proyek kecerdasan buatan terarah rapi dari perencanaan sampai jadi aplikasi siap pakai."),

        ("CDSS (Clinical Decision Support System)",
         "Sistem Pendukung Keputusan Klinis",
         "Perangkat lunak interaktif yang dirancang untuk membantu tenaga medis atau individu dalam menganalisis data klinis guna mendukung pengambilan keputusan preventif atau diagnostik.",
         "Aplikasi asisten cerdas bagi dokter yang memberikan saran analisis risiko dan alasan medisnya, tetapi keputusan vonis akhir tetap di tangan manusia."),

        ("Modifiable Lifestyle Factors",
         "Faktor Gaya Hidup yang Dapat Diubah",
         "Variabel perilaku sehari-hari yang berada di bawah kendali kehendak individu untuk diperbaiki, seperti durasi menatap layar ponsel, jam olahraga, dan konsumsi kafein.",
         "Kebiasaan buruk yang sebenarnya bisa kita ubah sendiri jika kita sadar, berbeda dengan faktor takdir seperti tanggal lahir atau genetik keluarga.")
    ]

    # Create Glossary Table
    table_g = doc.add_table(rows=len(glossary_data) + 1, cols=4)
    table_g.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_g)
    
    col_w = [Inches(0.4), Inches(1.8), Inches(2.2), Inches(1.87)]
    for row in table_g.rows:
        for idx, width in enumerate(col_w):
            row.cells[idx].width = width

    g_headers = ["No", "Istilah Teknis", "Definisi Ilmiah Akademis", "Bahasa Manusia (Analogi Dosen)"]
    for c_idx, h_text in enumerate(g_headers):
        cell = table_g.rows[0].cells[c_idx]
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        set_cell_shading(cell, "1F4E79")
        p = cell.paragraphs[0]
        if c_idx == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        set_font(run, size=9.5, bold=True, color=RGBColor(255,255,255))

    for idx, (term, alias, formal, simple) in enumerate(glossary_data, start=1):
        row = table_g.rows[idx]
        bg_col = "F9FBFD" if idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([str(idx), f"{term}\n({alias})", formal, simple]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            set_cell_shading(cell, bg_col)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(p.add_run(val), size=9, bold=True)
            elif c_idx == 1:
                set_font(p.add_run(val), size=9, bold=True, color=RGBColor(31, 78, 121))
            elif c_idx == 2:
                set_font(p.add_run(val), size=8.5)
            else:
                set_font(p.add_run(val), size=8.5, italic=True)

    # -------------------------------------------------------------
    # BAGIAN 4: CHECKLIST FISIK & TIPS MENGHADAP DOSEN
    # -------------------------------------------------------------
    doc.add_page_break()
    h4 = doc.add_heading(level=1)
    set_font(h4.add_run("BAGIAN 4: CHECKLIST FISIK & TIPS MENGHADAP DOSEN"), size=12.5, bold=True, color=RGBColor(31, 78, 121))

    tips_text = (
        "1. Berkas Fisik yang Wajib Dibawa dalam Map Rapi:\n"
        "   - Cetak Formulir Outline Proposal Resmi (OUTLINE_PROPOSAL_REVISI_FINAL.docx).\n"
        "   - Pastikan halaman Tabel Research GAP (5 Jurnal SOTA) dan Matriks Novelty mudah dibuka langsung.\n"
        "   - Bawa laptop/gawai yang siap membuka purwarupa Web CDSS jika dosen penasaran ingin melihat demonya.\n\n"
        "2. Prinsip 'Show, Don't Tell' (Tunjukkan Buktinya, Jangan Cuma Bicara):\n"
        "   - Saat dosen menanyakan kebaruan, langsung buka halaman Tabel Research GAP dan tunjukkan nama Das et al. (2025) di jurnal Nature and Science of Sleep.\n"
        "   - Saat dosen menanyakan data, sebutkan angka pasti: '100.000 data terbagi 80:20 stratified'. Dosen sangat menyukai angka yang spesifik dan pasti.\n\n"
        "3. Sikap Saat Terjadi Perdebatan atau Dosen Menyela:\n"
        "   - Jangan pernah memotong kalimat dosen. Dengarkan sampai selesai, anggukkan kepala sebagai tanda menghargai.\n"
        "   - Awali respons dengan: 'Terima kasih atas masukannya, Bapak/Ibu. Izin menjelaskan pertimbangan teknis di balik keputusan tersebut...'\n"
        "   - Jika dosen menyarankan perubahan kecil pada redaksi kata judul, terimalah dengan lapang dada: 'Baik Bapak/Ibu, saran penyempurnaan redaksinya sangat baik dan akan segera saya sesuaikan.'"
    )
    add_callout("TIPS PSIKOLOGIS & CHECKLIST SEBELUM MENGHADAP KAPRODI/DOSEN", tips_text, border_hex="2E7D32", bg_hex="F1F8F1")

    # Output file
    out_file = "PANDUAN_PENGAJUAN_JUDUL_DAN_TANYA_JAWAB_SKRIPSI.docx"
    doc.save(out_file)
    print(f"[OK] Guide docx generated successfully at: {out_file}")

if __name__ == "__main__":
    create_guide_docx()

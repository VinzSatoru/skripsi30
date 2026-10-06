import docx
import re
import os
import fitz # PyMuPDF

import shutil
shutil.copy2("PROPOSAL BAB 1-3.docx", "temp_inspect_proposal.docx")
with open("temp_inspect_proposal.docx", "rb") as f:
    doc = docx.Document(f)

# 1. Extract text by section
b1_text = ""
b2_text = ""
b3_text = ""
dp_text = ""

current_bab = 0
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("BAB I"):
        current_bab = 1
    elif t.startswith("BAB II"):
        current_bab = 2
    elif t.startswith("BAB III"):
        current_bab = 3
    elif "DAFTAR PUSTAKA" in t.upper():
        current_bab = 4
    
    if current_bab == 1:
        b1_text += t + "\n"
    elif current_bab == 2:
        b2_text += t + "\n"
    elif current_bab == 3:
        b3_text += t + "\n"
    elif current_bab == 4:
        dp_text += t + "\n"

# Also check tables
for t_idx, table in enumerate(doc.tables):
    for row in table.rows:
        for cell in row.cells:
            if t_idx == 0:
                b2_text += cell.text + "\n"
            else:
                b3_text += cell.text + "\n"

all_body_text = b1_text + "\n" + b2_text + "\n" + b3_text

# 2. Extract all entries in DAFTAR PUSTAKA
start_dp = -1
for i, p in enumerate(doc.paragraphs):
    if "DAFTAR PUSTAKA" in p.text.upper():
        start_dp = i
        break

dp_entries = [p.text.strip() for p in doc.paragraphs[start_dp + 1:] if p.text.strip()]

print(f"Total DAFTAR PUSTAKA entries: {len(dp_entries)}")

# 3. For each entry in DAFTAR PUSTAKA, extract author and year, and check if it is cited in body
print("\n=== CHECKING EACH DAFTAR PUSTAKA ENTRY IN BODY TEXT ===")
unmatched_in_body = []
for idx, entry in enumerate(dp_entries, 1):
    # Extract first author surname and year
    m = re.match(r'^([A-Z][a-zA-Z\s\-]+?),\s*([A-Z]\..*?)\s*\((\d{4})\)', entry)
    has_doi = "doi.org" in entry.lower() or "https://" in entry.lower()
    
    if m:
        first_author = m.group(1).strip()
        year = m.group(3).strip()
    else:
        # Fallback split
        parts = entry.split("(")
        first_author = parts[0].split(",")[0].strip()
        year = parts[1].split(")")[0].strip() if len(parts) > 1 else ""

    # Check if first_author is in all_body_text
    # Search for patterns like "FirstAuthor et al. (Year)", "FirstAuthor & ... (Year)", "(FirstAuthor ..., Year)"
    # Or just first_author
    in_b1 = first_author.lower() in b1_text.lower()
    in_b2 = first_author.lower() in b2_text.lower()
    in_b3 = first_author.lower() in b3_text.lower()
    in_any = in_b1 or in_b2 or in_b3

    status = "CITED" if in_any else "NOT CITED (ORPHAN!)"
    doi_status = "HAS DOI" if has_doi else "NO DOI!"
    
    where = []
    if in_b1: where.append("Bab 1")
    if in_b2: where.append("Bab 2")
    if in_b3: where.append("Bab 3")
    where_str = ", ".join(where) if where else "NONE"

    print(f"{idx:2d}. [{status}] [{doi_status}] {first_author} ({year}) -> in {where_str}")
    if not in_any:
        unmatched_in_body.append((idx, first_author, year, entry))
    if not has_doi:
        print(f"     ^^ NO DOI: {entry[:80]}...")

print(f"\nTotal DP entries NOT cited in body: {len(unmatched_in_body)}")
for idx, a, y, e in unmatched_in_body:
    print(f"  - #{idx}: {a} ({y}) -> {e[:100]}...")

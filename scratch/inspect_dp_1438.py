import docx

doc = docx.Document("BAB 1_1438.docx")
start_dp = -1
for i, p in enumerate(doc.paragraphs):
    if "DAFTAR PUSTAKA" in p.text.upper():
        start_dp = i
        break

print(f"DAFTAR PUSTAKA starts at index {start_dp}")
for i in range(start_dp, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"[{i}] {p.text}")
